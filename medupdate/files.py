"""Hochgeladene Dateien lesen, Volltexte beschaffen und Uploads Publikationen zuordnen.
Großteils aus MyLuo übernommen."""

from __future__ import annotations

import io
import re
import unicodedata

import fitz  # PyMuPDF
import requests
from docx import Document
from pypdf import PdfReader

from .http import BROWSER_HEADERS, request_with_retry


# ---------- Text aus Dateien ----------

def extract_docx_text(file_bytes: bytes) -> str:
    """Text einer .docx-Datei, Absatz für Absatz, inklusive Tabelleninhalten."""
    doc = Document(io.BytesIO(file_bytes))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(lines)


def extract_pdf_text(file_bytes: bytes, max_chars: int | None = None) -> str:
    """Text einer PDF. PyMuPDF liefert meist saubereren Text als pypdf; pypdf dient als Rückfall."""
    text = ""
    try:
        with fitz.open(stream=file_bytes, filetype="pdf") as doc:
            text = "\n".join(page.get_text() for page in doc)
    except Exception:
        text = ""
    if len(text.strip()) < 200:
        try:
            reader = PdfReader(io.BytesIO(file_bytes))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception:
            pass
    text = text.strip()
    return text[:max_chars] if max_chars else text


def extract_text(filename: str, file_bytes: bytes, max_chars: int | None = None) -> str:
    name = filename.lower()
    if name.endswith(".docx"):
        text = extract_docx_text(file_bytes)
    elif name.endswith(".pdf"):
        text = extract_pdf_text(file_bytes)
    else:
        text = file_bytes.decode("utf-8", errors="replace")
    return text[:max_chars] if max_chars else text


def extract_images_from_pdf(pdf_bytes: bytes, max_images: int = 40) -> list[bytes]:
    """Eingebettete Abbildungen (Fotos, Diagramme, als Bild eingefügte Tabellen) einer PDF.
    Rein vektoriell gezeichnete Tabellen/Grafiken werden dabei nicht erfasst."""
    images = []
    try:
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        for page in doc:
            for img in page.get_images(full=True):
                try:
                    base_image = doc.extract_image(img[0])
                    if len(base_image["image"]) > 3000 and base_image.get("width", 0) >= 200:
                        images.append(base_image["image"])
                except Exception:
                    continue
                if len(images) >= max_images:
                    return images
    except Exception:
        return images
    return images


# ---------- Open Access ----------

def unpaywall(doi: str, email: str) -> dict:
    if not doi or not email:
        return {}
    try:
        r = request_with_retry("GET", f"https://api.unpaywall.org/v2/{doi}", params={"email": email}, timeout=15)
        return r.json()
    except (requests.exceptions.RequestException, ValueError):
        return {}


def download_pdf(url: str) -> bytes | None:
    try:
        r = request_with_retry("GET", url, timeout=60, headers=BROWSER_HEADERS, retries=2)
    except requests.exceptions.RequestException:
        return None
    return r.content if r.content[:4] == b"%PDF" else None


def fetch_oa_pdf(doi: str, email: str) -> bytes | None:
    """Versucht, die frei verfügbare PDF einer Publikation zu laden (alle OA-Fundorte von
    Unpaywall nacheinander). None, wenn nichts gefunden wird oder der Verlag blockiert."""
    info = unpaywall(doi, email)
    if not info.get("is_oa"):
        return None
    locations = [info.get("best_oa_location") or {}] + list(info.get("oa_locations") or [])
    seen = set()
    for loc in locations:
        url = loc.get("url_for_pdf")
        if url and url not in seen:
            seen.add(url)
            pdf = download_pdf(url)
            if pdf:
                return pdf
    return None


def looks_like_real_fulltext(candidate: str, abstract: str) -> bool:
    """Erkennt, ob ein vermeintlicher Volltext mehr als nur das Abstract enthält (manche Verlage
    liefern nur Vorschau- oder 'Article in Press'-Platzhalter-PDFs)."""
    if not candidate or len(candidate) < 2000:
        return False
    if len(candidate) < len(abstract or "") * 1.5:
        return False
    preview = candidate[:300].lower()
    return "article in press" not in preview


# ---------- Uploads zuordnen ----------

def _fold(s: str) -> str:
    """Ligaturen (ﬁ → fi) und Akzente auflösen, klein schreiben."""
    decomposed = unicodedata.normalize("NFKD", s)
    return "".join(c for c in decomposed if not unicodedata.combining(c)).lower()


def _normalize(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", _fold(s))


def upload_matches(extracted_text: str, rec: dict) -> bool:
    """Heuristik aus MyLuo: Passt eine PDF zur erwarteten Publikation? DOI im Text, oder Titelanfang
    bzw. ≥90 % der Titelwörter im Kopfbereich plus Erstautor."""
    head = extracted_text[:6000]
    if rec.get("doi") and rec["doi"].lower() in head.lower():
        return True
    norm_text = _normalize(head)
    norm_title = _normalize(rec.get("title", ""))
    snippet = norm_title[:40]
    title_match = bool(snippet) and snippet in norm_text
    if not title_match:
        words = [w for w in re.findall(r"[a-z0-9]+", _fold(rec.get("title", ""))) if len(w) >= 4]
        if len(words) >= 4:
            head_words = set(re.findall(r"[a-z0-9]+", _fold(head)))
            title_match = sum(w in head_words for w in words) / len(words) >= 0.9
    authors = rec.get("authors") or []
    first = authors[0].split(" ")[0] if authors else ""
    author_match = bool(first) and _normalize(first) in norm_text
    return title_match and (author_match or not first)


_DOI_RE = re.compile(r"\b(10\.\d{4,9}/[^\s\"<>]+)", re.IGNORECASE)


def find_doi(text: str) -> str:
    m = _DOI_RE.search(text[:8000])
    return m.group(1).rstrip(".,;)") if m else ""
