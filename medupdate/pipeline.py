"""Ablaufsteuerung. Alle Zwischenergebnisse landen im Projekt-Dict `proj`; jeder Schritt
überspringt bereits Erledigtes. Bricht ein Durchlauf ab (z.B. Verbindung weg), setzt ein
erneuter Start genau dort fort."""

from __future__ import annotations

import base64
import gzip
import json
import re
import threading
import traceback
import uuid
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, timezone

from . import appraisal as A
from . import guidelines as G
from . import intake, pubmed, recommend
from . import prompts_new as P
from .citations import (
    dedupe_local_literature,
    extract_and_renumber_citations_global,
    reconcile_local_citations,
    strip_appraisal_preamble,
)
from .config import APP_VERSION, MAX_PUBS_PER_TOPIC, MODEL_FAST, MODEL_MAIN, PARALLEL_WORKERS
from .files import extract_pdf_text, fetch_oa_pdf, find_doi, looks_like_real_fulltext, \
    unpaywall, upload_matches
from .formal import run_formal_check
from .llm import LLM, STR, AccountError, arr, enum, obj
from .manuscript import build_docx, build_recommendation_docx
from .docx_tools import build_pruefprotokoll_docx
from .references import make_formatter



# Version der Kürzungsvorgabe; 2 = kürzere Sätze, weniger Statistik (2026-10-10); 3 = Methodik verständlich
# erklärt, abwechslungsreiches Fazit (2026-10-11)
SHORT_VERSION = 8
MAX_GROUP_SIZE = 5  # gemeinsame Würdigung von bis zu 5 Publikationen (Angelika 2026-10-13)
GROUP_VERSION = 3  # höher setzen, wenn sich die Regeln für gemeinsame Würdigungen ändern

def new_project() -> dict:
    return {"version": 1, "id": uuid.uuid4().hex[:12], "created": date.today().isoformat(), "email": "", "toc": None, "rules": None,
            "template": None, "topics": {}, "records": {}, "pubs": {}, "uploads": {}, "style": {},
            "guideline_uploads": {}, "outputs": {}, "errors": []}


def _err(proj: dict, where: str, ex: Exception) -> None:
    proj["errors"].append({"where": where, "error": f"{type(ex).__name__}: {ex}"})
    traceback.print_exc()


class _Skipped(Exception):
    """Aufgabe wurde nach einem Konto-/Kostenlimit-Stopp nicht mehr gestartet."""


class _Streamed:
    """Ergebnis von _parallel: .items() liefert jedes Ergebnis, SOBALD es fertig ist. Die Aufrufer
    speichern es so sofort im Projekt – bricht der Lauf später ab, ist bereits Bezahltes nicht verloren."""

    def __init__(self, items, fn, log, label, workers):
        self.args = (list(items), fn, log, label, workers)

    def items(self):
        items, fn, log, label, workers = self.args
        if not items:
            return
        stop: Exception | None = None
        halt = threading.Event()

        def guarded(it):
            if halt.is_set():
                raise _Skipped()
            try:
                return fn(it)
            except AccountError:
                halt.set()  # sofort im Arbeits-Thread, damit keine weitere Anfrage mehr startet
                raise

        with ThreadPoolExecutor(max_workers=workers) as ex:
            futs = {ex.submit(guarded, it): key for key, it in items}
            for n, fut in enumerate(as_completed(futs), start=1):
                key = futs[fut]
                if fut.cancelled():
                    continue
                try:
                    res = fut.result()
                except _Skipped:
                    continue
                except Exception as e:  # einzelne Fehler sollen den Rest nicht stoppen
                    res = e
                log(f"{label}: {n}/{len(futs)} fertig")
                if isinstance(res, AccountError):
                    # Kontoproblem oder Kostenlimit: nichts Neues mehr starten, bereits laufende Aufrufe
                    # (schon bezahlt) aber noch abholen und speichern
                    stop = stop or res
                    for f in futs:
                        f.cancel()
                    continue
                yield key, res
        if stop:
            raise stop


def _parallel(items, fn, log, label: str, workers: int = PARALLEL_WORKERS) -> _Streamed:
    return _Streamed(items, fn, log, label, workers)


# ======================================================================
# SCHRITT 1
# ======================================================================

def run_intake(llm: LLM, proj: dict, toc_text: str, guide_text: str, template: tuple[str, bytes] | None,
               log) -> None:
    if not proj.get("toc"):
        log("Inhaltsverzeichnis wird ausgewertet…")
        proj["toc"] = intake.parse_toc(llm, toc_text)
        proj["topics"] = {t["id"]: {"id": t["id"], **t} for t in proj["toc"]["topics"]}
    if not proj.get("rules"):
        log("Hinweise für Autoren werden ausgewertet…")
        proj["rules"] = intake.parse_author_rules(llm, guide_text) if guide_text.strip() else {}
    if template and not proj.get("template"):
        name, data = template
        proj["template"] = {"name": name, "bytes": data, **intake.inspect_template(data)}


# Toleranz für die manuell eingegebenen Obergrenzen (Seiten, Literaturstellen)
LIMIT_TOLERANCE = {"mode": "absolute", "value": 10}


def effective_limit(value: int, tol: dict | None = None) -> int:
    tol = tol or LIMIT_TOLERANCE
    if not value:
        return 0
    if tol.get("mode") == "absolute":
        return int(value + tol.get("value", 0))
    return int(value * (1 + tol.get("value", 0) / 100.0))


def apply_manual_limits(proj: dict) -> None:
    """Manuell eingegebene Vorgaben (Seiten ohne Inhaltsverzeichnis, Literaturstellen) haben Vorrang vor
    den Hinweisen für Autoren. Die Grenze ist der eingegebene Wert plus Toleranz; für die Seitenplanung kommt
    die Titelseite mit dem Inhaltsverzeichnis hinzu."""
    lim = proj.get("manual_limits") or {}
    if proj.get("rules") is None:
        proj["rules"] = {}
    rules = proj["rules"]
    tol = proj.get("limit_tolerance") or LIMIT_TOLERANCE
    guide = rules.setdefault("_guide_limits", {"max_pages": rules.get("max_pages") or 0,
                                               "max_references": rules.get("max_references") or 0})
    rules["max_pages"] = effective_limit(lim["pages"], tol) + 1 if lim.get("pages") else guide["max_pages"]
    rules["max_references"] = (effective_limit(lim["references"], tol) if lim.get("references")
                               else guide["max_references"])
    # Höchstzahl gewürdigter Publikationen: genau der eingegebene Wert (keine Toleranz)
    rules["max_publications"] = int(lim.get("publications") or 0)


NO_TOLERANCE = {"mode": "absolute", "value": 0}


def set_strict_limits(proj: dict, strict: bool) -> None:
    """Häkchen „Keine Toleranz“ (Angelika 2026-10-14): Seiten und Literaturstellen sind dann harte
    Obergrenzen wie die Publikationszahl."""
    proj["limit_tolerance"] = dict(NO_TOLERANCE) if strict else dict(LIMIT_TOLERANCE)
    apply_manual_limits(proj)


def tolerance_note(proj: dict) -> str:
    tol = proj.get("limit_tolerance") or LIMIT_TOLERANCE
    return ("keine Überschreitung erlaubt" if not tol.get("value")
            else f"Überschreitung um höchstens {tolerance_text(proj)}")


def tolerance_text(proj: dict) -> str:
    tol = proj.get("limit_tolerance") or LIMIT_TOLERANCE
    if not tol.get("value"):
        return "0 (keine Toleranz)"
    if tol.get("mode") != "absolute":
        return f"{tol['value']} %"
    return f"{tol['value']} Seiten bzw. {tol['value']} Literaturstellen"


def run_recommendations(llm: LLM, proj: dict, log) -> None:
    apply_manual_limits(proj)
    proj["check_ok"] = False  # neue Auswahl: Themenzuordnung in Schritt 2 erneut prüfen
    chapter = (proj["toc"] or {}).get("chapter_title", "")
    rules = proj.get("rules") or {}
    n = rules.get("pubs_per_topic") or MAX_PUBS_PER_TOPIC
    window = proj.get("search_window")  # manuell eingestellt; bei Änderung werden die Themen neu gesucht
    todo = [(tid, t) for tid, t in proj["topics"].items() if t.get("enabled", True) and not t.get("varia")
            and ("recommended" not in t or t.get("n_candidates") == 0 or (window and t.get("window") != window))]

    def work(t):
        return recommend.recommend_for_topic(llm, t, chapter, rules, proj["email"], n, window=window)

    for tid, res in _parallel(todo, work, log, "Literaturempfehlung").items():
        if isinstance(res, Exception):
            _err(proj, f"Literaturempfehlung {proj['topics'][tid]['title']}", res)
            proj["topics"][tid]["rec_error"] = f"{type(res).__name__}: {res}"
            continue
        proj["topics"][tid].pop("rec_error", None)
        proj["topics"][tid].pop("selected", None)  # alte Auswahl passt nicht zur neuen Trefferliste
        proj["records"].update(res.pop("records"))
        proj["topics"][tid].update(res)

    # Open-Access-Status der empfohlenen Arbeiten
    for t in proj["topics"].values():
        for pmid in t.get("recommended", []) + t.get("reserve", []):
            rec = proj["records"].get(pmid)
            if rec is not None and "is_oa" not in rec:
                info = unpaywall(rec.get("doi", ""), proj["email"])
                rec["is_oa"] = bool(info.get("is_oa"))
    fit_to_length(proj, log)


def _enabled_topics(proj: dict) -> list[str]:
    return [tid for tid, t in proj["topics"].items() if t.get("enabled", True)]


def _rank_item(proj: dict, p: str, t: str) -> dict:
    """Bewertung aus Schritt 1; für eine gemeinsame Würdigung die beste der beteiligten Publikationen."""
    rk = proj["topics"][t].get("ranking", {})
    members = (proj["pubs"].get(p) or {}).get("group")
    if members:
        items = [rk.get(m, {}) for m in members]
        return max(items, key=lambda it: it.get("length_score", 0), default={})
    return rk.get(p, {})


def _length_info(proj: dict, pairs: list[tuple[str, str]]) -> list[dict]:
    return [{"id": p, "length_score": _rank_item(proj, p, t).get("length_score", 50)} for p, t in pairs]


def _n_sections(proj: dict, topic_ids) -> int:
    return len({proj["topics"][t].get("section_number") or t for t in topic_ids}) or 1


def fit_to_length(proj: dict, log=lambda m: None) -> None:
    """Gesamtlänge laut Hinweisen für Autoren im Blick behalten: Längen vorplanen und, falls selbst
    mit ½ Seite je Arbeit das Limit überschritten wird, die schwächsten Empfehlungen in die Reserve
    verschieben (die Autorin kann sie wieder anhaken)."""
    pairs = [(p, tid) for tid in _enabled_topics(proj) for p in proj["topics"][tid].get("recommended", [])]
    tiers, overflow = A.plan_tiers(_length_info(proj, pairs), proj.get("rules") or {},
                                   _n_sections(proj, {t for _, t in pairs}))
    for p, tid in pairs:
        if p in overflow:
            t = proj["topics"][tid]
            t["recommended"].remove(p)
            t.setdefault("reserve", []).insert(0, p)
            if p in t.get("selected", []):
                t["selected"].remove(p)
    if overflow:
        log(f"{len(overflow)} Empfehlung(en) wegen der vorgegebenen Gesamtlänge in die Reserve verschoben.")
    proj["length_plan"] = {"tiers": tiers, "overflow": overflow,
                           "pages": A.estimate_pages(tiers, _n_sections(proj, {t for _, t in pairs})),
                           "max_pages": (proj.get("rules") or {}).get("max_pages") or 0}


def run_guideline_search(llm: LLM, proj: dict, log) -> None:
    chapter = (proj["toc"] or {}).get("chapter_title", "")
    todo = [(tid, t) for tid, t in proj["topics"].items() if t.get("enabled", True) and "guidelines" not in t]

    def work(t):
        found = G.search_guidelines(llm, t["title"], chapter)
        for g in found:
            text, src = G.fetch_guideline_text(g)
            g["text"], g["source"] = text, src
            g["status"] = "geladen" if text else "fehlt"
        return found

    for tid, res in _parallel(todo, work, log, "Leitlinienrecherche").items():
        if isinstance(res, Exception):
            _err(proj, f"Leitlinien {proj['topics'][tid]['title']}", res)
            proj["topics"][tid]["guidelines"] = []
        else:
            proj["topics"][tid]["guidelines"] = res


def recommendation_docx(proj: dict) -> bytes:
    rankings = {}
    for t in proj["topics"].values():
        rankings.update(t.get("ranking", {}))
    by_topic: dict[str, list[str]] = {}
    for p, tid in _pub_list(proj):  # aktuelle Auswahl, jede Publikation nur einmal
        by_topic.setdefault(tid, []).append(p)
    topics = [{"title": f"{t.get('number', '')} {t['title']}".strip(), "pmids": by_topic.get(tid, [])}
              for tid, t in proj["topics"].items() if t.get("enabled", True)]
    return build_recommendation_docx((proj["toc"] or {}).get("chapter_title", ""), topics, proj["records"], rankings)


# ======================================================================
# SCHRITT 2: Uploads zuordnen
# ======================================================================

IDENT_SCHEMA = obj({"title": STR, "authors": arr(STR), "journal": STR, "year": STR, "doi": STR,
                    "kind": enum("original", "review", "guideline", "other")})
ASSIGN_SCHEMA = obj({"assignments": arr(obj({"key": STR, "topic_id": STR}))})


def selected_pmids(proj: dict) -> list[str]:
    out = []
    for t in proj["topics"].values():
        if t.get("enabled", True):
            out.extend(t.get("selected", t.get("recommended", [])))
    return out


def match_uploads(llm: LLM, proj: dict, uploads: dict[str, bytes], log) -> None:
    """Ordnet hochgeladene PDFs den empfohlenen Publikationen zu. Nicht zuordenbare Dateien werden
    identifiziert (DOI → PubMed, sonst KI) und als zusätzliche Publikationen einem Thema zugeordnet."""
    wanted = {p: proj["records"][p] for p in selected_pmids(proj) if p in proj["records"]}
    unmatched = []
    for name, data in uploads.items():
        if name in proj["uploads"]:  # bereits zugeordnet: Text und Zuordnung bleiben
            continue
        text = extract_pdf_text(data)
        hit = next((p for p, rec in wanted.items() if upload_matches(text, rec)), None)
        proj["uploads"][name] = {"pmid": hit, "text_len": len(text)}
        if hit:
            proj["pubs"].setdefault(hit, {})["full_text"] = text
        else:
            unmatched.append((name, data, text))

    extra = []
    for name, data, text in unmatched:
        log(f"Zusätzliche Datei „{name}“ wird identifiziert…")
        doi = find_doi(text)
        pmid = pubmed.find_pmid(doi=doi, email=proj["email"]) if doi else None
        if not pmid:
            try:
                meta = llm.json(P.identify_pub_prompt(text[:5000]), IDENT_SCHEMA, max_tokens=2000, effort="low")
            except Exception as ex:
                _err(proj, f"Identifikation {name}", ex)
                continue
            if meta.get("kind") == "guideline":
                proj["uploads"][name]["role"] = "background"
                proj["guideline_uploads"][name] = {"topic_id": "", "text": text}
                continue
            pmid = pubmed.find_pmid(doi=meta.get("doi", ""), title=meta.get("title", ""), email=proj["email"])
            if not pmid:
                pmid = f"upload:{name}"
                proj["records"][pmid] = {"pmid": "", "title": meta["title"], "authors": meta["authors"],
                                         "journal": meta["journal"], "year": meta["year"], "doi": meta["doi"],
                                         "abstract": "", "pub_types": [], "volume": "", "issue": "", "pages": ""}
        if pmid not in proj["records"]:
            proj["records"].update(pubmed.fetch_records([pmid], proj["email"]))
        proj["uploads"][name]["pmid"] = pmid
        proj["pubs"].setdefault(pmid, {}).update({"full_text": text, "extra": True})
        extra.append(pmid)

    if extra:
        topics_block = "\n".join(f"{t['id']}: {t['title']}" for t in proj["topics"].values() if t.get("enabled", True))
        pubs_block = "\n".join(
            f"{p}: {proj['records'].get(p, {}).get('title', '')} – "
            + " ".join((proj["records"].get(p, {}).get("abstract") or proj["pubs"][p].get("full_text", "")[:600]).split())[:500]
            for p in extra)
        try:
            data = llm.json(P.topic_assign_prompt(topics_block, pubs_block), ASSIGN_SCHEMA, max_tokens=3000,
                            effort="low")
            for a in data.get("assignments", []):
                if a["key"] in proj["pubs"] and a["topic_id"] in proj["topics"]:
                    proj["pubs"][a["key"]]["topic_id"] = a["topic_id"]
        except AccountError:
            raise
        except Exception as ex:
            _err(proj, "Themenzuordnung", ex)
        ensure_varia(proj)


def attach_direct_upload(proj: dict, p: str, name: str, data: bytes, force: bool = False) -> dict:
    """PDF direkt bei einer Publikation der Literaturliste hochgeladen: wie in MyLuo kurz prüfen, ob es die
    richtige Publikation ist (DOI oder Titel plus Erstautor im Kopfbereich). Ohne KI, daher kostenlos.
    Passt sie nicht, wird sie nur mit force=True übernommen. Rückgabe: {"ok", "match", "found"}."""
    text = extract_pdf_text(data)
    rec = proj["records"].get(p, {})
    match = upload_matches(text, rec)
    found = next((ln.strip() for ln in text.splitlines() if len(ln.strip()) > 25), "")[:160]
    if not (match or force):
        return {"ok": False, "match": False, "found": found, "empty": not text.strip()}
    d = proj["pubs"].setdefault(p, {})
    d.update({"full_text": text, "fulltext_source": "Upload", "upload_name": name,
              "upload_checked": "passt" if match else "von Ihnen bestätigt"})
    d.pop("images", None)
    proj["uploads"][name] = {"pmid": p, "text_len": len(text)}
    return {"ok": True, "match": match, "found": found, "empty": not text.strip()}


VARIA = "VARIA"


def ensure_varia(proj: dict) -> None:
    """Hochgeladene Publikationen, die zu keinem Thema des Inhaltsverzeichnisses passen, kommen unter
    den Inhaltspunkt „Varia“ (eigener, letzter Abschnitt ohne Literaturrecherche und ohne State of the Art)."""
    orphans = [p for p, d in proj["pubs"].items() if d.get("extra") and not d.get("group")
               and (d.get("topic_id") or "") not in proj["topics"]]
    if not orphans:
        return
    if VARIA not in proj["topics"]:
        proj["topics"][VARIA] = {"id": VARIA, "title": "Varia", "number": "", "english_search_focus": "",
                                 "section_number": "Varia", "section_title": "Varia", "enabled": True,
                                 "varia": True, "recommended": [], "reserve": [], "ranking": {}, "guidelines": [],
                                 "window": ["", ""], "n_candidates": 0}
    for p in orphans:
        proj["pubs"][p]["topic_id"] = VARIA


REASSIGN_SCHEMA = obj({"assignments": arr(obj({"pub_id": STR, "topic_id": STR, "reason_de": STR})),
                       "same_question": arr(obj({"pub_ids": arr(STR), "question_de": STR}))})


def current_topics(proj: dict) -> dict[str, str]:
    """{pmid: topic_id} aller Publikationen, die ins Manuskript kommen (ohne gemeinsame Würdigungen)."""
    return {p: t for p, t in _pub_list(proj) if not proj["pubs"].get(p, {}).get("group")}


def bump_selection(proj: dict) -> None:
    """Die Auswahl wurde vom Programm geändert: die Ankreuzfelder in Schritt 1 übernehmen den neuen Stand."""
    proj["sel_rev"] = proj.get("sel_rev", 0) + 1


def _unselect_everywhere(proj: dict, p: str, keep: str = "") -> None:
    """Dieselbe Arbeit kann in Schritt 1 für mehrere Themen empfohlen worden sein."""
    for tid, t in proj["topics"].items():
        if tid != keep and p in t.get("selected", t.get("recommended", [])):
            t["selected"] = [x for x in t.get("selected", t.get("recommended", [])) if x != p]


def move_pub(proj: dict, p: str, dst: str) -> None:
    """Verschiebt eine Publikation in ein anderes Thema (die Bewertung aus Schritt 1 wandert mit)."""
    src = current_topics(proj).get(p)
    if not src or src == dst or dst not in proj["topics"]:
        return
    d = proj["pubs"].setdefault(p, {})
    d.setdefault("home_topic", src)
    bump_selection(proj)
    if d.get("extra"):
        d["topic_id"] = dst
        return
    s_t, d_t = proj["topics"][src], proj["topics"][dst]
    _unselect_everywhere(proj, p, keep=dst)
    d_t["selected"] = list(d_t.get("selected", d_t.get("recommended", []))) + [p]
    if p in s_t.get("ranking", {}):
        d_t.setdefault("ranking", {})[p] = s_t["ranking"][p]


def _reassign_key(proj: dict) -> list:
    return ([f"v{GROUP_VERSION}"] + sorted(current_topics(proj))
            + sorted(f"{p}={t}" for p, t in (proj.get("topic_manual") or {}).items()))


def reassign_topics(llm: LLM, proj: dict, log, force: bool = False) -> None:
    """Prüft nach dem Hochladen die Themenzuordnung ALLER Publikationen im Gesamtblick: Schwerpunkt der
    Arbeit entscheidet (klinisch → Krankheitsbild, technischer Aspekt → Verfahren); Arbeiten mit nahezu
    gleicher Fragestellung kommen ins selbe Thema und werden als Gruppe vorgemerkt."""
    ensure_varia(proj)
    _restore_limit_dropped(proj)  # wegen der Vorgaben entfernte Arbeiten zählen mit; enforce_pub_limit folgt
    if not force and proj.get("reassigned_for") == _reassign_key(proj):
        return
    cur = current_topics(proj)
    if len(cur) < 2:
        proj["reassigned_for"] = _reassign_key(proj)
        return
    log("Themenzuordnung aller Publikationen wird geprüft…")
    manual = proj.get("topic_manual") or {}
    topics = [(tid, t) for tid, t in proj["topics"].items() if t.get("enabled", True) and not t.get("varia")]
    topics_block = "\n".join(
        f"{tid}: {t.get('number', '')} {t['title']}"
        + (f" (Abschnitt {t.get('section_number', '')} {t.get('section_title', '')})" if t.get("section_title") else "")
        for tid, t in topics)
    lines = []
    for p, tid in cur.items():
        rec, d = proj["records"].get(p, {}), proj["pubs"].get(p, {})
        q = proj["topics"][tid].get("ranking", {}).get(p, {}).get("question_key", "")
        summary = " ".join((rec.get("abstract") or d.get("full_text", "")[:900]).split())[:600]
        lines.append(f"- {p}{' [hochgeladen]' if d.get('extra') else ''}{' [vom Nutzer festgelegt]' if p in manual else ''}"
                     f" | {'' if tid == VARIA else tid} | {rec.get('title', p)} | {q} | {summary}")
    res = llm.json(P.reassign_prompt((proj["toc"] or {}).get("chapter_title", ""), topics_block, "\n".join(lines)),
                   REASSIGN_SCHEMA, model=MODEL_MAIN, max_tokens=6000, effort="medium")
    valid = {tid for tid, _ in topics}
    target = {}
    for a in res.get("assignments", []):
        p, tid = a.get("pub_id", ""), a.get("topic_id", "")
        if p not in cur or p in manual:
            continue
        if tid in valid:
            target[p] = (tid, a.get("reason_de", "").strip())
        elif not tid and proj["pubs"].get(p, {}).get("extra"):
            target[p] = (VARIA, a.get("reason_de", "").strip())
    groups, used = [], set()
    for g in res.get("same_question", []):
        ids = [p for p in dict.fromkeys(g.get("pub_ids", [])) if p in cur and p not in used]
        if 2 <= len(ids) <= MAX_GROUP_SIZE:
            used.update(ids)
            groups.append({"pmids": ids, "question_de": g.get("question_de", "").strip()})
    # Gleiche Fragestellung → gleiches Thema: das Thema der ersten (bzw. vom Nutzer festgelegten) Publikation
    for g in groups:
        fixed = next((manual[p] for p in g["pmids"] if p in manual), None)
        tid = fixed or target.get(g["pmids"][0], (cur[g["pmids"][0]], ""))[0]
        for p in g["pmids"]:
            if p not in manual and target.get(p, (cur[p],))[0] != tid:
                target[p] = (tid, f"Gleiche Fragestellung wie „{g['question_de']}“; gemeinsame Würdigung.")
    moved = 0
    for p, (tid, reason) in target.items():
        if tid != cur[p]:
            if tid == VARIA and VARIA not in proj["topics"]:
                proj["pubs"][p].setdefault("home_topic", cur[p])
                proj["pubs"][p]["topic_id"] = ""
                ensure_varia(proj)
            else:
                move_pub(proj, p, tid)
            proj["pubs"][p]["move_reason"] = reason
            moved += 1
    for p, tid in manual.items():
        move_pub(proj, p, tid)
    proj["question_groups"] = groups
    proj["reassigned_for"] = _reassign_key(proj)
    log(f"Themenzuordnung geprüft: {moved} Publikation(en) verschoben, {len(groups)} Gruppe(n) mit gleicher "
        "Fragestellung.")


def rate_unranked(llm: LLM, proj: dict, log) -> None:
    """Zusätzlich hochgeladene Publikationen haben noch keine Qualitätsbewertung aus Schritt 1: Bewertung
    mit demselben Ranking (Studienqualität, Bedeutung der Fragestellung, Impact Factor) nachholen."""
    by_topic: dict[str, list[str]] = {}
    for p, tid in current_topics(proj).items():
        if p not in proj["topics"][tid].get("ranking", {}) and not proj["pubs"].get(p, {}).get("rank_tried"):
            by_topic.setdefault(tid, []).append(p)
    for tid, ids in by_topic.items():
        t = proj["topics"][tid]
        log(f"Qualität von {len(ids)} hochgeladenen Publikation(en) in „{t['title']}“ wird bewertet…")
        recs = []
        for p in ids:
            proj["pubs"].setdefault(p, {})["rank_tried"] = True
            rec, d = proj["records"].get(p, {}), proj["pubs"].get(p, {})
            recs.append({"pmid": p, "title": rec.get("title", p), "journal": rec.get("journal", ""),
                         "year": rec.get("year", ""), "pub_types": rec.get("pub_types") or [],
                         "abstract": rec.get("abstract") or _main_text(d.get("full_text", ""))[:3500]})
        try:
            ranking = recommend.rank(llm, t, "", recs, "des gewählten Zeitraums")
        except AccountError:
            raise
        except Exception as ex:
            _err(proj, f"Qualitätsbewertung {t['title']}", ex)
            continue
        for p, item in ranking.items():
            item["score"] = recommend.score(item)
            item["priority"] = recommend.priority_class(item)
            item["quality"] = recommend.quality(item)
            item["length_score"] = recommend.length_score(item)
            t.setdefault("ranking", {})[p] = item


def _restore_limit_dropped(proj: dict) -> None:
    for e in proj.get("limit_dropped", []):
        p, tid = e["id"], e["topic_id"]
        d = proj["pubs"].get(p, {})
        if d.get("extra"):
            d.pop("limit_dropped", None)
            continue
        t = proj["topics"].get(tid)
        if not t:
            continue
        if e.get("to_reserve"):
            t["reserve"] = [x for x in t.get("reserve", []) if x != p]
        sel = list(t.get("selected", t.get("recommended", [])))
        if p not in sel:
            t["selected"] = sel + [p]
    if proj.get("limit_dropped"):
        bump_selection(proj)
    proj["limit_dropped"] = []


def enforce_pub_limit(proj: dict, log) -> list[dict]:
    """Nach Hochladen und Sortieren: Höchstzahl der Publikationen und Seitenumfang einhalten. Ist eines davon
    überschritten, werden die Studien mit der geringsten Qualität (Qualitätsranking aus Schritt 1) entfernt –
    bevor Kosten für Würdigungen entstehen. Wird neu berechnet, sobald sich Auswahl oder Vorgaben ändern."""
    apply_manual_limits(proj)
    _restore_limit_dropped(proj)
    rules = proj.get("rules") or {}
    max_n, max_pages = rules.get("max_publications") or 0, rules.get("max_pages") or 0
    cur = current_topics(proj)
    if not cur or not (max_n or max_pages):
        return []

    def rank(p):
        it = proj["topics"][cur[p]].get("ranking", {}).get(p, {})
        return (it.get("quality", recommend.quality(it) if it else 0), it.get("length_score", 0))

    def units(ids):
        # Publikationen mit gleicher Fragestellung werden gemeinsam gewürdigt und brauchen nur einmal Platz
        out, seen = [], set()
        for g in proj.get("question_groups", []):
            m = [p for p in g["pmids"] if p in ids and cur.get(p) == cur.get(g["pmids"][0])]
            if len(m) >= 2:
                seen.update(m)
                out.append({"id": "+".join(m), "length_score": max(rank(p)[1] for p in m)})
        return out + [{"id": p, "length_score": rank(p)[1]} for p in ids if p not in seen]

    def too_long(ids):
        if not max_pages:
            return False
        _tiers, overflow = A.plan_tiers(units(ids), rules, _n_sections(proj, {cur[p] for p in ids}))
        return bool(overflow)

    keep = sorted(cur, key=rank, reverse=True)
    dropped = []
    while len(keep) > 1 and ((max_n and len(keep) > max_n) or too_long(keep)):
        p = keep.pop()
        reason = (f"Höchstzahl von {max_n} Publikationen" if max_n and len(keep) >= max_n
                  else f"Seitenumfang (max. {max_pages - 1} Seiten)")
        dropped.append({"id": p, "topic_id": cur[p], "reason": reason,
                        "title": proj["records"].get(p, {}).get("title", p)})
    for e in dropped:
        p, t = e["id"], proj["topics"][e["topic_id"]]
        d = proj["pubs"].setdefault(p, {})
        if d.get("extra"):
            d["limit_dropped"] = True
        else:
            _unselect_everywhere(proj, p)
            if p not in t.get("recommended", []) and p not in t.get("reserve", []):
                t.setdefault("reserve", []).insert(0, p)
                e["to_reserve"] = True
    proj["limit_dropped"] = dropped
    if dropped:
        bump_selection(proj)
        log(f"{len(dropped)} Publikation(en) mit der geringsten Studienqualität wegen der Vorgaben entfernt.")
    return dropped


def check_assignment(llm: LLM, proj: dict, uploads: dict[str, bytes], log) -> None:
    """Pflichtschritt vor dem Schreiben: Uploads zuordnen, Themenzuordnung prüfen, hochgeladene Arbeiten
    bewerten, Publikationszahl und Seitenumfang einhalten. Erst danach wird „Manuskript erstellen“ frei."""
    proj["check_ok"] = False
    if uploads:
        match_uploads(llm, proj, uploads, log)
    reassign_topics(llm, proj, log)
    rate_unranked(llm, proj, log)
    enforce_pub_limit(proj, log)
    proj["check_ok"] = True


def set_manual_topic(proj: dict, p: str, tid: str) -> None:
    """Zuordnung durch die Nutzerin; die KI-Prüfung ändert sie danach nicht mehr."""
    proj.setdefault("topic_manual", {})[p] = tid
    move_pub(proj, p, tid)
    proj["pubs"][p]["move_reason"] = "Von Ihnen festgelegt."
    proj["reassigned_for"] = _reassign_key(proj)


def assignment_overview(proj: dict) -> list[dict]:
    """Überschriften des Inhaltsverzeichnisses mit den zugeordneten Publikationen (für Anzeige und Word-Liste)."""
    cur = current_topics(proj)
    group_of = {p: g for g in proj.get("question_groups", []) for p in g["pmids"]}
    out, last_sec = [], None
    for tid, t in proj["topics"].items():
        if not t.get("enabled", True):
            continue
        sec = (t.get("section_number", ""), t.get("section_title", ""))
        if sec != last_sec and sec[1] and sec[1] != t["title"]:
            out.append({"kind": "section", "text": f"{sec[0]} {sec[1]}".strip()})
        last_sec = sec
        pubs = []
        for p in (p for p, x in cur.items() if x == tid):
            rec, d = proj["records"].get(p, {}), proj["pubs"].get(p, {})
            home = d.get("home_topic")
            g = group_of.get(p)
            pubs.append({"id": p, "title": rec.get("title", p), "journal": rec.get("journal", ""),
                         "year": rec.get("year", ""), "uploaded": bool(d.get("extra")),
                         "moved_from": proj["topics"][home]["title"] if home and home != tid and home in proj["topics"] else "",
                         "reason": d.get("move_reason", "") if home and home != tid else "",
                         "joint": g["question_de"] if g and all(cur.get(m) == tid for m in g["pmids"]) else ""})
        out.append({"kind": "topic", "id": tid, "text": f"{t.get('number', '')} {t['title']}".strip(), "pubs": pubs})
    return out


def assignment_docx(proj: dict) -> bytes:
    from io import BytesIO
    from docx import Document
    doc = Document()
    doc.add_heading(f"Themenzuordnung – {(proj['toc'] or {}).get('chapter_title', '')}", level=1)
    for row in assignment_overview(proj):
        if row["kind"] == "section":
            doc.add_heading(row["text"], level=2)
            continue
        doc.add_heading(row["text"], level=3)
        if not row["pubs"]:
            doc.add_paragraph().add_run("Keine Publikation (entfällt im Manuskript).").italic = True
        for x in row["pubs"]:
            para = doc.add_paragraph(style="List Bullet")
            para.add_run(f"{x['title']} — {x['journal']} {x['year']}".strip())
            notes = (["hochgeladen"] if x["uploaded"] else []) + \
                ([f"verschoben aus „{x['moved_from']}“: {x['reason']}".rstrip(": ")] if x["moved_from"] else []) + \
                ([f"gemeinsame Würdigung: {x['joint']}"] if x["joint"] else [])
            if notes:
                para.add_run(" (" + "; ".join(notes) + ")").italic = True
    if proj.get("limit_dropped"):
        doc.add_heading("Wegen der Vorgaben nicht aufgenommen", level=2)
        for e in proj["limit_dropped"]:
            doc.add_paragraph(f"{e['title']} ({proj['topics'].get(e['topic_id'], {}).get('title', '')}; {e['reason']})",
                              style="List Bullet")
    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()


# ======================================================================
# SCHRITT 2: Automatische Erstellung
# ======================================================================

def _pub_list(proj: dict) -> list[tuple[str, str]]:
    """[(pmid, topic_id)] in Manuskript-Reihenfolge: Themen wie im Inhaltsverzeichnis, darin nach Score."""
    out, seen = [], set()
    for tid, t in proj["topics"].items():
        if not t.get("enabled", True):
            continue
        # Eine Publikation steht nur einmal im Manuskript (z.B. nach Verschieben in ein anderes Thema)
        ids = [p for p in t.get("selected", t.get("recommended", [])) if p not in seen]
        ids += [p for p, d in proj["pubs"].items() if d.get("extra") and d.get("topic_id") == tid and p not in ids
                and not d.get("limit_dropped")]
        rk = t.get("ranking", {})
        ids.sort(key=lambda p: -rk.get(p, {}).get("score", 0))
        seen.update(ids)
        out.extend((p, tid) for p in ids)
    return out


def _guideline_docs(proj: dict, tid: str, query: str, per_doc: int = 6000, total: int = 14000) -> list:
    docs = []
    t = proj["topics"][tid]
    for g in t.get("guidelines", []):
        if g.get("text"):
            label = f"Leitlinie: {g.get('title', '')} ({g.get('society', '')}, {g.get('year', '')})"
            docs.append((label, G.relevant_excerpt(g["text"], query, per_doc)))
    for name, g in proj.get("guideline_uploads", {}).items():
        if g.get("topic_id") in ("", tid) and g.get("text"):
            docs.append((f"Vom Autor bereitgestellt: {name}", G.relevant_excerpt(g["text"], query, per_doc)))
    out, used = [], 0
    for name, text in docs:
        if used + len(text) > total:
            break
        out.append((name, text))
        used += len(text)
    return out


def _entry(rec: dict, pub: dict) -> dict:
    return {"title": rec.get("title", ""), "journal": rec.get("journal", ""), "pubdate": rec.get("year", ""),
            "pub_types": rec.get("pub_types") or ["(unbekannt)"], "abstract": rec.get("abstract") or "(kein Abstract)",
            "full_text": _main_text(pub.get("full_text") or "")[:35000]}


def _main_text(text: str) -> str:
    """Volltext ohne Literaturverzeichnis und Anhang (spart Token, enthält keine Ergebnisse)."""
    m = re.search(r"\n\s*(References|REFERENCES|Literatur|Bibliography)\s*\n", text[len(text) // 3:])
    return text[: len(text) // 3 + m.start()] if m else text


def run_manuscript(llm: LLM, proj: dict, log, options: dict) -> None:
    email = proj["email"]
    # Ein älteres Manuskript (z.B. aus einem geladenen Zwischenstand) darf nach einem Abbruch nicht als
    # Ergebnis dieses Laufs heruntergeladen werden.
    for k in ("manuskript_DE", "manuskript_EN", "rohfassung_DE", "rohfassung_EN", "pruefprotokoll", "befunde",
              "created", "version"):
        (proj.get("outputs") or {}).pop(k, None)
    apply_manual_limits(proj)
    ensure_varia(proj)
    try:
        reassign_topics(llm, proj, log)  # ist nach check_assignment bereits erledigt (Cache)
    except AccountError:
        raise
    except Exception as ex:
        _err(proj, "Themenzuordnung prüfen", ex)
    rate_unranked(llm, proj, log)
    enforce_pub_limit(proj, log)
    pubs = _pub_list(proj)
    if options.get("trial"):
        # Probelauf: gewähltes (sonst erstes) Thema mit Publikationen, höchstens 2 Arbeiten – kostet nur
        # einen Bruchteil
        wanted = options.get("trial_topic")
        first = wanted if any(t == wanted for _, t in pubs) else next((t for _, t in pubs), None)
        pubs = [(p, t) for p, t in pubs if t == first][:2]
    log(f"{len(pubs)} Publikationen in {len({t for _, t in pubs})} Themen werden verarbeitet.")

    # 1. Volltexte (Upload oder Open Access)
    for pmid, _tid in pubs:
        d = proj["pubs"].setdefault(pmid, {})
        rec = proj["records"].get(pmid, {})
        if not d.get("full_text") and not d.get("oa_tried"):
            d["oa_tried"] = True
            pdf = fetch_oa_pdf(rec.get("doi", ""), email)
            if pdf:
                text = extract_pdf_text(pdf)
                if looks_like_real_fulltext(text, rec.get("abstract", "")):
                    d["full_text"], d["fulltext_source"] = text, "Open Access"
        # Abbildungen werden nicht mehr automatisch ausgewählt (2026-10-11): PDF nach der Textextraktion
        # nicht mehr im Speicher halten
        d.pop("pdf", None)
        d.pop("images", None)

    # 2. Kritische Würdigung (englisch, wie in MyLuo als Arbeitsfassung)
    todo = [(p, (p, t)) for p, t in pubs if not proj["pubs"][p].get("appraisal_en")]

    def appraise(item):
        p, tid = item
        rec = proj["records"].get(p, {})
        query = f"{proj['topics'][tid]['title']} {rec.get('title', '')} {rec.get('abstract', '')[:600]}"
        return A.generate_appraisal(llm, _entry(rec, proj["pubs"][p]), _guideline_docs(proj, tid, query), "Englisch")

    log("Kritische Würdigungen werden erstellt (mit Websuche, das dauert mehrere Minuten pro Publikation)…")
    for p, res in _parallel(todo, appraise, log, "Kritische Würdigung").items():
        if isinstance(res, Exception):
            _err(proj, f"Würdigung {proj['records'].get(p, {}).get('title', p)[:60]}", res)
        else:
            proj["pubs"][p]["appraisal_en"] = res

    ready = [(p, t) for p, t in pubs if proj["pubs"][p].get("appraisal_en")]

    # 3. Lokale Literaturliste (Nr. 1 = Publikation selbst) und Längenstufen
    for p, _ in ready:
        d = proj["pubs"][p]
        if "body_en" not in d:
            body, entries = extract_and_renumber_citations_global(d["appraisal_en"], 2)
            main = pubmed.format_reference(proj["records"].get(p, {}))
            body, lit = dedupe_local_literature(body, [main] + entries)
            d["body_en"], d["literature"] = strip_appraisal_preamble(body), lit
    # 3b. Gemeinsame Würdigung für Publikationen mit nahezu gleicher Fragestellung (wie MyLuo, Schritt 4B)
    if options.get("joint", True):
        ready = apply_groups(llm, proj, ready, log)

    limit_key = [(proj.get("rules") or {}).get("max_pages"), (proj.get("rules") or {}).get("max_references")]
    if (not proj.get("tiers") or options.get("replan_tiers") or any(p not in proj["tiers"] for p, _ in ready)
            or proj.get("tiers_for_limits") != limit_key):
        proj["tiers_for_limits"] = limit_key
        # Selbst ausgewählte Arbeiten werden nicht verworfen: im Notfall alle auf ½ Seite, Hinweis im Protokoll
        tiers, overflow = A.plan_tiers(_length_info(proj, ready), proj.get("rules") or {},
                                       _n_sections(proj, {t for _, t in ready}))
        for p in overflow:
            tiers[p] = "½ Seite (Kurz)"
        proj["tiers"] = tiers
    for p, _ in ready:
        proj["tiers"].setdefault(p, "½ Seite (Kurz)")

    # 4. Kürzung (englisch). Auf Wunsch werden ältere Kürzungen (vor SHORT_VERSION) neu erstellt; dann
    # auch Übersetzung und Stilfassung, die darauf aufbauen.
    if options.get("reshorten"):
        for p, _ in ready:
            d = proj["pubs"][p]
            if d.get("short_en") and d.get("short_v", 1) < SHORT_VERSION:
                for k in ("short_en", "short_de", "styled_de", "style_ok"):
                    d.pop(k, None)
    todo = [(p, p) for p, _ in ready if not proj["pubs"][p].get("short_en")]

    # Literaturbudget: Hauptquellen + State of the Art (≈3 je Abschnitt) zuerst, der Rest verteilt sich auf
    # die externen Verweise der Würdigungen (mindestens einer je Würdigung)
    max_refs = (proj.get("rules") or {}).get("max_references") or 0
    ext_budget = None
    if max_refs:
        n_main = sum(len(proj["pubs"][p].get("group") or [p]) for p, _ in ready)
        n_sota = _n_sections(proj, {t for _, t in ready}) * 3
        ext_budget = max(1, (max_refs - n_main - n_sota) // max(1, len(ready)))

    def shorten(p):
        d, rec = proj["pubs"][p], proj["records"].get(p, {})
        if d.get("group"):
            return A.shorten_joint(llm, d["body_en"], proj["tiers"][p], "Englisch", max_ext=ext_budget)
        return A.shorten(llm, d["body_en"], proj["tiers"][p], "Englisch", rec.get("abstract", ""),
                         max_ext=ext_budget)

    log("Würdigungen werden auf die geplanten Längen gekürzt…")
    for p, res in _parallel(todo, shorten, log, "Kürzung").items():
        if isinstance(res, Exception):
            _err(proj, f"Kürzung {p}", res)
            proj["pubs"][p]["short_en"] = proj["pubs"][p]["body_en"]  # ungekürzt weiterverwenden
        else:
            proj["pubs"][p]["short_en"], proj["pubs"][p]["short_v"] = res, SHORT_VERSION
            proj["pubs"][p]["short_tier"] = proj["tiers"][p]

    # 5. Übersetzung ins Deutsche
    todo = [(p, p) for p, _ in ready if not proj["pubs"][p].get("short_de")]

    def to_de(p):
        d, rec = proj["pubs"][p], proj["records"].get(p, {})
        out = {"short_de": A.translate_structured(llm, d["short_en"])}
        if proj["tiers"][p] == "2 Seiten" and rec.get("abstract"):
            out["abstract_de"] = A.translate_abstract(llm, rec["abstract"])
        return out

    log("Übersetzung ins Deutsche…")
    for p, res in _parallel(todo, to_de, log, "Übersetzung").items():
        if isinstance(res, Exception):
            _err(proj, f"Übersetzung {p}", res)
        else:
            proj["pubs"][p].update(res)

    # 6. State of the Art: einer je Hauptabschnitt, vor dem Thema mit der hochwertigsten Publikation
    # (bei Gleichstand: mit den meisten Publikationen)
    topic_ids = [t for t in proj["topics"] if any(tt == t for _, tt in ready)]
    counts = {t: sum(len(proj["pubs"][p].get("group") or [p]) for p, tt in ready if tt == t) for t in topic_ids}
    best_q = {t: max((recommend.quality(_rank_item(proj, p, t)) for p, tt in ready if tt == t),
                     default=0.0) for t in topic_ids}
    sota_ids = [t for t in A.sota_topics(topic_ids, proj["topics"], counts, best_q)
                if not proj["topics"][t].get("varia")]
    proj["sota_topics"] = sota_ids

    def first_pub(tid):
        """Die erste Publikation des Themas mit Volltext (sie folgt direkt auf den State of the Art);
        bei einer gemeinsamen Würdigung die erste beteiligte Publikation mit Volltext."""
        for p, tt in ready:
            if tt != tid:
                continue
            for m in proj["pubs"][p].get("group") or [p]:
                if proj["pubs"].get(m, {}).get("full_text"):
                    return m
        return None

    # Neu erstellen, wenn noch keiner da ist oder ein älterer (per Websuche) durch einen aus der
    # Einleitung der folgenden Publikation ersetzt werden kann
    todo = [(t, t) for t in sota_ids if not proj["topics"][t].get("sota_en")
            or (proj["topics"][t]["sota_en"].get("source") != A.SOTA_SOURCE and first_pub(t))]

    def sota(tid):
        title = proj["topics"][tid]["title"]
        en = None
        if first_pub(tid):
            pubs = []  # alle Publikationen des Themas mit Volltext, in Manuskript-Reihenfolge
            for q, tt in ready:
                if tt != tid:
                    continue
                for m in proj["pubs"][q].get("group") or [q]:
                    if proj["pubs"].get(m, {}).get("full_text") and m not in [x for x, _ in pubs]:
                        pubs.append((m, proj["pubs"][m]["full_text"]))
            en = A.generate_sota_from_intro(
                llm, title, [(proj["records"].get(m, {}).get("title", ""), ft) for m, ft in pubs],
                _guideline_docs(proj, tid, title, per_doc=4000, total=8000))
        if en is None:
            en = A.generate_sota(llm, title, "Englisch", _guideline_docs(proj, tid, title, per_doc=5000, total=12000))
        return {"sota_en": en, "sota_de": A.translate_sota(llm, en)}

    if options.get("sota", True):
        log(f"State of the Art wird erstellt ({len(sota_ids)} Hauptabschnitt(e))…")
        for t, res in _parallel(todo, sota, log, "State of the Art").items():
            if isinstance(res, Exception):
                _err(proj, f"State of the Art {proj['topics'][t]['title']}", res)
            else:
                proj["topics"][t].update(res)

    # 7. Stilanpassung (deutsch) an frühere Manuskripte
    samples = list(proj.get("style", {}).values())
    if samples and options.get("style", True):
        if not proj.get("style_profile"):
            log("Schreibstil der früheren Manuskripte wird analysiert…")
            proj["style_profile"] = A.style_profile(llm, samples)
        todo = [(p, p) for p, _ in ready if proj["pubs"][p].get("short_de") and "styled_de" not in proj["pubs"][p]]
        excerpt = samples[0][1000:3500]

        def styl(p):
            return A.apply_style(llm, proj["pubs"][p]["short_de"], proj["style_profile"], excerpt)

        log("Stilanpassung…")
        for p, res in _parallel(todo, styl, log, "Stilanpassung").items():
            if isinstance(res, Exception):
                _err(proj, f"Stilanpassung {p}", res)
            else:
                proj["pubs"][p]["styled_de"], proj["pubs"][p]["style_ok"] = res

    # 7a. Fazit-Anfänge abwechslungsreich (nicht überall "Das Ergebnis ist neu.")
    try:
        vary_fazit_openings(llm, proj, ready, log)
    except AccountError:
        raise
    except Exception as ex:
        _err(proj, "Fazit-Formulierungen", ex)

    # 7b. Inhaltliche Rückbezüge auf frühere eigene Manuskripte (höchstens OWN_REFS_MAX, nur wo sinnvoll)
    if proj.get("style") and options.get("own_refs", True):
        key = sorted(p for p, _ in ready)
        if proj.get("own_refs_for") != key:
            log("Rückbezüge auf frühere Manuskripte werden gesucht…")
            try:
                proj["own_refs"] = find_own_refs(llm, proj, ready)
                proj["own_refs_for"] = key
            except AccountError:
                raise
            except Exception as ex:
                _err(proj, "Rückbezüge auf frühere Manuskripte", ex)

    # 9. Manuskripte zusammensetzen, Literatur zusammenführen, Formatvorlage, formale Prüfung
    log("Gesamtmanuskript wird zusammengesetzt und formal geprüft…")
    rules = proj.get("rules") or {}
    title = (proj["toc"] or {}).get("chapter_title") or "Gesamtmanuskript"
    template = (proj.get("template") or {}).get("bytes")
    formatter = make_formatter(llm, rules, proj["records"], email)
    gliederung = outline(proj, topic_ids)
    try:
        de = headings_de(llm, proj, topic_ids)
    except AccountError:
        raise
    except Exception as ex:
        _err(proj, "Übersetzung der Überschriften", ex)
        de = {}
    for lang in ("Deutsch", "Englisch"):
        blocks = []
        tr = (lambda x: de.get(x, x)) if lang == "Deutsch" else (lambda x: x)
        for tid in topic_ids:
            t = proj["topics"][tid]
            sota_block = (t.get("sota_de" if lang == "Deutsch" else "sota_en")
                          if tid in sota_ids and options.get("sota", True) else None)
            g = dict(gliederung.get(tid, {}))
            for k in ("section", "subsection"):
                if g.get(k):
                    g[k] = (g[k][0], tr(g[k][1]))
            block = {"topic": tr(t["title"]), "sota": sota_block, "pubs": [], **g}
            for p, tt in ready:
                if tt != tid:
                    continue
                d, rec = proj["pubs"][p], proj["records"].get(p, {})
                if lang == "Deutsch":
                    text = d.get("styled_de") or d.get("short_de") or d["short_en"]
                    abstract = d.get("abstract_de")
                else:
                    text = d["short_en"]
                    abstract = rec.get("abstract") if proj["tiers"].get(p) == "2 Seiten" else None
                own = (proj.get("own_refs") or {}).get(p) if options.get("own_refs", True) else None
                if own:
                    text = insert_own_ref(text, own["de" if lang == "Deutsch" else "en"])
                members = d.get("group") or []
                text, lit = reconcile_local_citations(text, d["literature"], protected_count=max(1, len(members)))
                if members:
                    rec = dict(rec, title=d["title_de" if lang == "Deutsch" else "title_en"])
                block["pubs"].append({"title": rec.get("title", p), "rec": rec, "abstract": abstract,
                                      "recs": [proj["records"].get(m, {}) for m in members],
                                      "lines": [_strip_stray_label(l) for l in text.split("\n") if l.strip()],
                                      "literature": lit,
                                      "figure": None})
            blocks.append(block)
        raw = build_docx(tr(title), blocks, lang)
        key = "DE" if lang == "Deutsch" else "EN"
        try:
            befunde, final = run_formal_check(raw, llm if lang == "Deutsch" else None, template, rules, formatter)
        except Exception as ex:
            _err(proj, f"Formale Prüfung {key}", ex)
            befunde, final = [], raw
        proj["outputs"][f"manuskript_{key}"] = final
        proj["outputs"]["created"] = datetime.now(timezone.utc).isoformat(timespec="minutes")
        proj["outputs"]["version"] = APP_VERSION
        proj["outputs"][f"rohfassung_{key}"] = raw
        if lang == "Deutsch":
            befunde = _pipeline_findings(proj, ready) + befunde
            proj["outputs"]["befunde"] = befunde
            proj["outputs"]["pruefprotokoll"] = build_pruefprotokoll_docx(befunde)
    log("Fertig.")


_STRAY_LABEL = re.compile(r"^\s*(\*\*)?(State of the Art|Stand der Forschung|Hintergrund|Background)\s*:\s*(\*\*)?\s*",
                          re.I)


def _strip_stray_label(line: str) -> str:
    """Ein von der KI erfundenes Label wie „State of the art:“ mitten in einer Würdigung entfernen; der
    State of the Art steht nur einmal je Hauptabschnitt vor dem Thema (Wunsch der Autorin, 2026-10-13)."""
    return _STRAY_LABEL.sub("", line)


def _pipeline_findings(proj: dict, ready: list) -> list[dict]:
    """Zusätzliche Hinweise für das Prüfprotokoll: fehlende Volltexte, Leitlinien, Stilanpassung, Fehler."""
    out = []
    no_ft = [proj["records"].get(p, {}).get("title", p)[:70] for p, _ in ready if not proj["pubs"][p].get("full_text")]
    out.append({"punkt": "Volltexte", "status": "✅" if not no_ft else "⚠️",
                "begruendung": "Alle Würdigungen beruhen auf dem Volltext." if not no_ft else
                "Nur auf Basis des Abstracts gewürdigt: " + "; ".join(no_ft)})
    missing = []
    for t in proj["topics"].values():
        for g in t.get("guidelines", []):
            if not g.get("text") and g.get("status") != "ersetzt":
                missing.append(f"{g.get('title', '')} ({g.get('society', '')})")
    out.append({"punkt": "Leitlinien", "status": "✅" if not missing else "⚠️",
                "begruendung": "Alle gefundenen Leitlinien lagen im Volltext vor." if not missing else
                "Nicht im Volltext verfügbar (nur über die Websuche berücksichtigt): " + "; ".join(missing[:10])})
    max_pages = (proj.get("rules") or {}).get("max_pages") or 0
    tiers = {p: proj["tiers"][p] for p, _ in ready if p in proj.get("tiers", {})}
    est = A.estimate_pages(tiers, _n_sections(proj, {t for _, t in ready}))
    counts = {}
    for t in tiers.values():
        counts[t] = counts.get(t, 0) + 1
    over = bool(max_pages and est > max_pages)
    out.append({"punkt": "Längenplanung (nach Studienqualität, Bedeutung der Fragestellung, Impact Factor)",
                "status": "⚠️" if over else "✅",
                "begruendung": ", ".join(f"{n}× {t}" for t, n in counts.items()) + f"; geschätzt ≈{est} Seiten"
                + (f" bei max. {max_pages} Seiten – Limit überschritten, bitte Auswahl reduzieren." if over
                   else (f" bei max. {max_pages} Seiten." if max_pages else "."))})
    styled = [p for p, _ in ready if "style_ok" in proj["pubs"][p]]
    if styled:
        bad = [p for p in styled if not proj["pubs"][p]["style_ok"]]
        out.append({"punkt": "Stilanpassung", "status": "✅" if not bad else "⚠️",
                    "begruendung": f"{len(styled) - len(bad)} von {len(styled)} Publikationen angepasst." +
                    (" Bei den übrigen wurde die Anpassung verworfen, weil sich Zahlen oder Umfang verändert "
                     "hätten." if bad else "")})
    if proj["errors"]:
        out.append({"punkt": "Fehler während der Erstellung", "status": "⚠️",
                    "begruendung": "; ".join(f"{e['where']}: {e['error'][:120]}" for e in proj["errors"][-10:])})
    return out


# ======================================================================
# Zwischenstand speichern / laden (die App speichert nichts auf dem Server)
# ======================================================================

def _enc(o):
    if isinstance(o, bytes):
        return {"__b64__": base64.b64encode(o).decode()}
    if isinstance(o, set):
        return list(o)
    raise TypeError(type(o))


def _dec(d):
    if isinstance(d, dict) and set(d) == {"__b64__"}:
        return base64.b64decode(d["__b64__"])
    return d


def export_project(proj: dict) -> bytes:
    """Zwischenstand ohne PDFs und Kandidatenbilder (die Texte sind enthalten) – hält die Datei klein."""
    slim = dict(proj)
    slim["pubs"] = {p: {k: v for k, v in d.items() if k not in ("pdf", "images")} for p, d in proj["pubs"].items()}
    return gzip.compress(json.dumps(slim, default=_enc, ensure_ascii=False).encode("utf-8"), 6)


FAZIT_SCHEMA = obj({"items": arr(obj({"id": STR, "sentence": STR}))})
_FAZIT_RE = re.compile(r"^(Fazit für die Praxis:\s*)(.+?[.!?])(\s+|$)(.*)$", re.S)


def _de_text_key(d: dict) -> str:
    return "styled_de" if d.get("styled_de") else "short_de"


def vary_fazit_openings(llm: LLM, proj: dict, ready: list[tuple[str, str]], log) -> int:
    """Fangen mehrere Fazits gleich an (z.B. "Das Ergebnis ist neu."), wird nur deren erster Satz in einem
    günstigen KI-Aufruf abwechslungsreich und konkret umformuliert. Rückgabe: Anzahl geänderter Fazits."""
    found = {}
    for p, _ in ready:
        d = proj["pubs"][p]
        key = _de_text_key(d)
        lines = (d.get(key) or "").split("\n")
        for i, line in enumerate(lines):
            m = _FAZIT_RE.match(line.strip())
            if m:
                found[p] = (key, i, m)
                break

    def start(m):
        return " ".join(re.findall(r"\w+", m.group(2).lower())[:3])

    counts: dict[str, int] = {}
    for _, _, m in found.values():
        counts[start(m)] = counts.get(start(m), 0) + 1
    dup = [p for p, (_, _, m) in found.items() if counts[start(m)] > 1]
    if len(dup) < 2:
        return 0
    log(f"Fazit-Formulierungen werden abwechslungsreicher gestaltet ({len(dup)})…")
    items = "\n\n".join(f"id {k}: {found[p][2].group(2)} {found[p][2].group(4)}".strip() for k, p in enumerate(dup))
    res = llm.json(P.vary_fazit_prompt(items), FAZIT_SCHEMA, model=MODEL_FAST, max_tokens=6000, effort="medium")
    changed = 0
    for it in res.get("items", []):
        try:
            p = dup[int(re.sub(r"\D", "", it.get("id", "")) or -1)]
        except (ValueError, IndexError):
            continue
        sent = (it.get("sentence") or "").strip()
        if not sent:
            continue
        key, i, m = found[p]
        d = proj["pubs"][p]
        lines = d[key].split("\n")
        lines[i] = (m.group(1) + sent + (" " + m.group(4) if m.group(4) else "")).rstrip()
        d[key] = "\n".join(lines)
        changed += 1
    return changed


HEADINGS_SCHEMA = obj({"items": arr(obj({"id": STR, "de": STR}))})


HEADINGS_VERSION = 2  # 2: deutsche Überschrift nach dem tatsächlichen Inhalt der Publikationen


def headings_de(llm: LLM, proj: dict, topic_ids: list[str]) -> dict[str, str]:
    """Deutsche Überschriften (Kapitel, Abschnitte, Unterabschnitte, Themen) für das deutsche Manuskript,
    auch wenn das Inhaltsverzeichnis englisch hochgeladen wurde. Nicht wörtlich übersetzt, sondern als
    deutscher Fachbegriff für das, worum es in den besprochenen Publikationen tatsächlich geht (Wunsch der
    Autorin, 2026-10-13). Ein günstiger KI-Aufruf; neu nur, wenn sich Überschriften oder Publikationen ändern."""
    by_topic: dict[str, list[str]] = {}
    for p, tid in _pub_list(proj):
        if tid in topic_ids:
            by_topic.setdefault(tid, []).append((proj["records"].get(p, {}).get("title") or "")[:160])
    ctx: dict[str, dict] = {}  # Überschrift -> {"level", "focus", "pubs"}
    chapter = (proj.get("toc") or {}).get("chapter_title") or ""
    if chapter.strip():
        ctx[chapter] = {"level": "Kapitel", "focus": "", "pubs": []}
    for tid in topic_ids:
        t = proj["topics"][tid]
        pubs = by_topic.get(tid, [])
        for level, src, focus in (("Abschnitt (x.x)", t.get("section_title") or "", ""),
                                  ("Unterabschnitt", t.get("subsection_title") or "", ""),
                                  ("Thema", t["title"], t.get("english_search_focus") or "")):
            if not src.strip() or src == "Varia":
                continue
            c = ctx.setdefault(src, {"level": level, "focus": focus, "pubs": []})
            c["pubs"] += [x for x in pubs if x not in c["pubs"]]
    key = [f"v{HEADINGS_VERSION}"] + sorted(f"{k}|{'|'.join(sorted(v['pubs']))}" for k, v in ctx.items())
    cache = proj.get("headings_de") or {}
    if proj.get("headings_de_for") == key:
        return cache
    srcs = list(ctx)
    items = "\n\n".join(
        f"{i}: {src}\n   Ebene: {ctx[src]['level']}"
        + (f"\n   Suchschwerpunkt: {ctx[src]['focus']}" if ctx[src]["focus"] else "")
        + ("\n   Besprochene Publikationen:\n" + "\n".join(f"   - {x}" for x in ctx[src]["pubs"][:12])
           if ctx[src]["pubs"] else "")
        for i, src in enumerate(srcs))
    res = llm.json(P.headings_de_prompt(items), HEADINGS_SCHEMA, model=MODEL_FAST, max_tokens=4000,
                   effort="medium")
    new = {}
    for it in res.get("items", []):
        try:
            src = srcs[int(it.get("id", ""))]
        except (ValueError, IndexError):
            continue
        if it.get("de", "").strip():
            new[src] = it["de"].strip()
    proj["headings_de"], proj["headings_de_for"] = new, key
    return new


def outline(proj: dict, topic_ids: list[str]) -> dict[str, dict]:
    """Gliederung nach dem hochgeladenen Inhaltsverzeichnis, aber nur mit Themen, zu denen Publikationen
    im Manuskript stehen; Hauptabschnitte, Unterabschnitte und Themen werden fortlaufend neu nummeriert.
    Rückgabe je Thema: {"number": "8.1.2.1" | None, "section": ("8.1", Titel) | None,
    "subsection": ("8.1.2", Titel) | None}. "section"/"subsection" stehen nur beim ersten Thema eines
    neuen (Unter-)Abschnitts. Ein (Unter-)Abschnitt mit nur einem gleichnamigen Thema erscheint nur einmal."""
    chapter = str((proj.get("toc") or {}).get("chapter_number") or "").strip().rstrip(".")
    prefix = f"{chapter}." if chapter else ""
    # Hauptabschnitt -> Einträge (Unterabschnitt mit Themen oder Thema direkt), in TOC-Reihenfolge
    sections: dict[str, dict[str, list[str]]] = {}
    titles: dict[str, str] = {}
    sub_titles: dict[tuple[str, str], str] = {}
    for tid in proj["topics"]:
        if tid not in topic_ids:
            continue
        t = proj["topics"][tid]
        key = t.get("section_number") or t.get("section_title") or tid
        sub = t.get("subsection_number") or t.get("subsection_title") or ""
        child = f"sub:{sub}" if sub else f"topic:{tid}"
        sections.setdefault(key, {}).setdefault(child, []).append(tid)
        titles.setdefault(key, t.get("section_title") or t["title"])
        if sub:
            sub_titles.setdefault((key, sub), t.get("subsection_title") or sub)

    def same(a: str, b: str) -> bool:
        return a.strip().lower() == b.strip().lower()

    out = {}
    for s_idx, (key, children) in enumerate(sections.items(), start=1):
        sec_num = f"{prefix}{s_idx}"
        all_tids = [x for tids in children.values() for x in tids]
        if len(all_tids) == 1 and same(titles[key], proj["topics"][all_tids[0]]["title"]):
            out[all_tids[0]] = {"section": (sec_num, titles[key]), "subsection": None, "number": None, "level": 2}
            continue
        first = True
        for c_idx, (child, tids) in enumerate(children.items(), start=1):
            num = f"{sec_num}.{c_idx}"
            sec = (sec_num, titles[key]) if first else None
            first = False
            if child.startswith("topic:"):
                out[tids[0]] = {"section": sec, "subsection": None, "number": num, "level": 2}
                continue
            st = sub_titles[(key, child[4:])]
            if len(tids) == 1 and same(st, proj["topics"][tids[0]]["title"]):
                out[tids[0]] = {"section": sec, "subsection": None, "number": num, "level": 2}
                continue
            for t_idx, tid in enumerate(tids, start=1):
                out[tid] = {"section": sec if t_idx == 1 else None,
                            "subsection": (num, st) if t_idx == 1 else None,
                            "number": f"{num}.{t_idx}", "level": 3}
    return out


GROUP_SCHEMA = obj({"groups": arr(obj({"pub_ids": arr(STR), "question_de": STR, "question_en": STR}))})


def find_groups(llm: LLM, proj: dict, tid: str, pmids: list[str]) -> list[dict]:
    """Publikationen eines Themas mit gleicher Fragestellung (je Gruppe 2 bis MAX_GROUP_SIZE)."""
    rk = proj["topics"][tid].get("ranking", {})
    pubs = [{"id": p, "title": proj["records"].get(p, {}).get("title", p),
             "question": rk.get(p, {}).get("question_key", ""),
             "summary": " ".join((proj["records"].get(p, {}).get("abstract") or "").split())[:500]} for p in pmids]
    hints = [g["pmids"] for g in proj.get("question_groups", []) if all(p in pmids for p in g["pmids"])]
    res = llm.json(P.group_prompt(proj["topics"][tid]["title"], pubs, hints), GROUP_SCHEMA, model=MODEL_MAIN,
                   max_tokens=3000, effort="medium")
    out, used = [], set()
    for g in res.get("groups", []):
        ids = [p for p in dict.fromkeys(g.get("pub_ids", [])) if p in pmids and p not in used]
        if 2 <= len(ids) <= MAX_GROUP_SIZE:
            used.update(ids)
            out.append({"pmids": ids, "question_de": g.get("question_de", "").strip(),
                        "question_en": g.get("question_en", "").strip()})
    return out


def apply_groups(llm: LLM, proj: dict, ready: list[tuple[str, str]], log) -> list[tuple[str, str]]:
    """Ersetzt Publikationen mit nahezu gleicher Fragestellung durch EINE gemeinsame Würdigung
    (Eintrag "G-…" in proj["pubs"], gebildet aus den Einzelwürdigungen). Gibt die neue Liste zurück."""
    singles = [(p, t) for p, t in ready if not proj["pubs"][p].get("group")]
    by_topic: dict[str, list[str]] = {}
    for p, t in singles:
        by_topic.setdefault(t, []).append(p)
    for tid, pmids in by_topic.items():
        t = proj["topics"][tid]
        if len(pmids) < 2:
            continue
        key = [f"v{GROUP_VERSION}"] + sorted(pmids)
        if t.get("groups_for") != key:
            log(f"Gemeinsame Fragestellungen in „{t['title']}“ werden geprüft…")
            try:
                t["groups"], t["groups_for"] = find_groups(llm, proj, tid, pmids), key
            except AccountError:
                raise
            except Exception as ex:
                _err(proj, f"Gruppierung {t['title']}", ex)
    todo = []
    for tid in by_topic:
        for g in proj["topics"][tid].get("groups", []):
            if not all(p in by_topic[tid] for p in g["pmids"]):
                continue
            gid = "G-" + "-".join(g["pmids"])
            if not proj["pubs"].get(gid, {}).get("appraisal_en"):
                todo.append((gid, (gid, tid, g)))

    def joint(item):
        gid, tid, g = item
        members = []
        for m in g["pmids"]:
            rec, d = proj["records"].get(m, {}), proj["pubs"][m]
            lit = "\n".join(f"{i}. {e}" for i, e in enumerate(d.get("literature", [])[1:], start=2))
            members.append({"title": rec.get("title", m), "journal": rec.get("journal", ""), "year": rec.get("year", ""),
                            "pub_types": ", ".join(rec.get("pub_types") or []),
                            "appraisal": d["body_en"] + ("\n\nReferences\n" + lit if lit else "")})
        return A.joint_appraisal(llm, g["question_en"] or proj["topics"][tid]["title"], members)

    if todo:
        log(f"Gemeinsame Würdigungen werden erstellt ({len(todo)})…")
    for gid, res in _parallel(todo, joint, log, "Gemeinsame Würdigung").items():
        if isinstance(res, Exception):
            _err(proj, "Gemeinsame Würdigung", res)
            continue
        tid, g = next((tid, g) for tid in by_topic for g in proj["topics"][tid].get("groups", [])
                      if "G-" + "-".join(g["pmids"]) == gid)
        body, entries = extract_and_renumber_citations_global(res, len(g["pmids"]) + 1)
        mains = [pubmed.format_reference(proj["records"].get(m, {})) for m in g["pmids"]]
        proj["pubs"][gid] = {"group": g["pmids"], "topic_id": tid, "appraisal_en": res,
                             "body_en": strip_appraisal_preamble(body), "literature": mains + entries,
                             "title_de": f"Gemeinsame Würdigung: {g['question_de'] or g['question_en']}",
                             "title_en": f"Joint appraisal: {g['question_en'] or g['question_de']}"}
    # Reihenfolge: die gemeinsame Würdigung steht an der Stelle ihrer ersten Publikation
    current = {"G-" + "-".join(g["pmids"]) for tid in by_topic for g in proj["topics"][tid].get("groups", [])
               if all(p in by_topic[tid] for p in g["pmids"])}
    member_of = {m: gid for gid in current if proj["pubs"].get(gid, {}).get("body_en")
                 for m in proj["pubs"][gid]["group"]}
    out, seen = [], set()
    for p, t in ready:
        gid = member_of.get(p, p)
        if gid not in seen:
            seen.add(gid)
            out.append((gid, t))
    return out


OWN_REFS_MAX = 2
OWN_REFS_SCHEMA = obj({"refs": arr(obj({"pub_id": STR, "manuscript": STR, "sentence_de": STR,
                                        "sentence_en": STR}))})


def find_own_refs(llm: LLM, proj: dict, ready: list[tuple[str, str]]) -> dict[str, dict]:
    """Ein KI-Aufruf über alle früheren Manuskripte und alle Publikationen: höchstens OWN_REFS_MAX
    Rückbezüge {pub_id: {"manuscript", "de", "en"}}."""
    manuscripts = [(n, t[:15000]) for n, t in list(proj["style"].items())[:4]]
    pubs = [{"id": p, "topic": proj["topics"][t]["title"], "title": proj["records"].get(p, {}).get("title", p),
             "summary": " ".join((proj["pubs"][p].get("short_en") or "").split())[:500]} for p, t in ready]
    res = llm.json(P.own_refs_prompt(manuscripts, pubs, OWN_REFS_MAX), OWN_REFS_SCHEMA, model=MODEL_MAIN,
                   max_tokens=4000, effort="medium")
    out = {}
    for r in res.get("refs", []):
        if r.get("pub_id") in proj["pubs"] and r.get("sentence_de", "").strip() and len(out) < OWN_REFS_MAX:
            out[r["pub_id"]] = {"manuscript": r.get("manuscript", ""), "de": r["sentence_de"].strip(),
                                "en": r.get("sentence_en", "").strip()}
    return out


def insert_own_ref(text: str, sentence: str) -> str:
    """Setzt den Rückbezug als eigenen Satz vor das Fazit bzw. den Kommentar, sonst vor den letzten Satz."""
    if not sentence:
        return text
    lines = text.split("\n")
    idx = next((i for i, l in enumerate(lines)
                if re.match(r"^\**\s*(Kommentar|Fazit für die Praxis|Comment|Conclusion for Practice)\s*:", l.strip())), None)
    if idx is not None:
        j = max((i for i in range(idx) if lines[i].strip()), default=None)
        abstract_label = re.compile(r"^\**\s*(Studie|Ergebnisse|Schlussfolgerung der Autoren|Study|Results|"
                                    r"Authors'? Conclusion)\s*:", re.IGNORECASE)
        if j is not None and not abstract_label.match(lines[j].strip()):
            lines[j] = lines[j].rstrip() + " " + sentence
        else:
            lines.insert(idx, sentence)  # keine Würdigung davor: eigener Absatz vor dem Fazit
        return "\n".join(lines)
    j = max((i for i, l in enumerate(lines) if l.strip()), default=None)
    if j is None:
        return text
    parts = re.split(r"(?<=[.!?])\s+(?=[A-ZÄÖÜ])", lines[j].strip())
    if len(parts) > 1:
        lines[j] = " ".join(parts[:-1] + [sentence, parts[-1]])
    else:
        lines[j] = lines[j].rstrip() + " " + sentence
    return "\n".join(lines)


def partial_docx(proj: dict) -> bytes | None:
    """Alle bisher fertigen Würdigungen als Word-Datei (beste verfügbare Fassung), auch nach einem Abbruch."""
    blocks = []
    for tid, t in proj["topics"].items():
        pubs = []
        for p, d in proj["pubs"].items():
            if p in t.get("selected", t.get("recommended", [])) or (d.get("extra") and d.get("topic_id") == tid):
                text = d.get("styled_de") or d.get("short_de") or d.get("short_en") or d.get("body_en")
                if text:
                    rec = proj["records"].get(p, {})
                    pubs.append({"title": rec.get("title", p), "rec": rec, "abstract": None,
                                 "lines": [x for x in text.split("\n") if x.strip()],
                                 "literature": d.get("literature", []), "figure": None})
        if pubs:
            blocks.append({"topic": t["title"], "sota": t.get("sota_de"), "pubs": pubs})
    if not blocks:
        return None
    return build_docx("Bisherige Ergebnisse (unvollständig)", blocks, "Deutsch")


def import_project(data: bytes) -> dict:
    """Lädt einen Zwischenstand (komprimiert .json.gz oder ältere unkomprimierte .json)."""
    if data[:2] == b"\x1f\x8b":
        data = gzip.decompress(data)
    proj = json.loads(data.decode("utf-8"), object_hook=_dec)
    for d in (proj.get("pubs") or {}).values():  # ältere Zwischenstände: große Binärdaten verwerfen
        d.pop("images", None)
    base = new_project()
    base.update(proj)
    return base
