"""Gesamtmanuskript als Word-Datei. Aufbau wie in MyLuo (Überschrift 2 je Publikation, lokale
Literaturlisten, "State of the Art" vor der ersten Publikation eines Themas), damit die
Literatur-Zusammenführung und die Übertragung in die Formatvorlage unverändert funktionieren."""

from __future__ import annotations

import io
import re
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from .docx_tools import (KURZ_HEADINGS, KURZ_MARK, _SKILL_BOLD_LABELS, add_field, add_floating_picture,
                         render_line_with_bold_label, style_document)

_COMMENT_LABELS = ("kommentar:", "comment:")


def tame_exclamations(line: str, in_comment: bool, left: int) -> tuple[str, int]:
    """Ausrufezeichen nur im Kommentar und höchstens eines je Hauptabschnitt (Entscheidung der Autorin
    2026-10-15); alle anderen werden Punkte. Gibt (Zeile, verbleibende Ausrufezeichen) zurück."""
    if "!" not in line:
        return line, left
    if not in_comment or left <= 0:
        return line.replace("!", "."), left
    first = line.index("!")
    return line[: first + 1] + line[first + 1:].replace("!", "."), left - 1


def _label_of(line: str) -> str | None:
    clean = line.replace("**", "").strip("* ").strip().lower()
    return next((l.lower() for l in _SKILL_BOLD_LABELS if clean.startswith(l.lower())), None)


def short_citation(rec: dict, lit_numbers: list[int]) -> str:
    """'Erstautor et al., Journal Jahr (N)' – wie format_short_citation in MyLuo."""
    authors = rec.get("authors") or []
    first = authors[0].split(" ")[0] if authors else "?"
    nums = ",".join(str(n) for n in lit_numbers)
    return f"{first} et al., {rec.get('journal', '')} {rec.get('year', '')} ({nums})".strip()


def build_docx(title: str, blocks: list[dict], lang: str) -> bytes:
    """blocks: [{"topic": str, "number": str (optional, z.B. "8.1.2"), "section": (Nummer, Titel) oder None
                 (nur beim ersten Thema eines neuen Hauptabschnitts), "subsection": ebenso für die
                 Ebene darunter, "level": 2 oder 3 (Gliederungsebene des Themas), "sota": {"text", "literature"} | None,
                 "comment": gemeinsamer Kommentar zu den kurz gewürdigten Publikationen des Themas | None,
                 "pubs": [{"title", "rec", "lines", "abstract", "literature", "figure", "kurz": bool}]}]
    Kurznachrichten ("kurz") stehen am Ende des Themas unter der unterstrichenen Überschrift „Kurznachrichten“;
    ihre Titelzeile trägt die Marke KURZ_MARK und entfällt nach der Literatur-Zusammenführung.
    Ausrufezeichen: nur im Kommentar, höchstens eines je Hauptabschnitt.
    Mit Nummern entsteht auf Seite 1 ein Inhaltsverzeichnis (Hauptabschnitte und Themen).
    lang: "Deutsch" oder "Englisch" (für Labels)."""
    lit_label = "References" if lang == "Englisch" else "Literatur"
    doc = Document()
    style_document(doc)

    p = doc.add_paragraph()
    r = p.add_run(title or "Gesamtmanuskript")
    r.bold = True
    r.font.size = Pt(14)
    doc.add_paragraph("")
    toc_lines = []
    for b in blocks:
        if b.get("section"):
            toc_lines.append((1, *b["section"]))
        if b.get("subsection"):
            toc_lines.append((2, *b["subsection"]))
        if b.get("number"):
            toc_lines.append((b.get("level") or 2, b["number"], b["topic"]))
    if toc_lines:
        p = doc.add_paragraph()
        p.add_run("Inhaltsverzeichnis" if lang != "Englisch" else "Contents").bold = True
        for k, (level, num, text) in enumerate(toc_lines):
            if level == 1 and k:
                doc.add_paragraph("")  # Leerzeile zwischen den Hauptabschnitten
            tp = doc.add_paragraph(f"{num}\t{text}")
            if level > 1:
                tp.paragraph_format.left_indent = Pt(18 * (level - 1))
        doc.add_paragraph("")
    doc.add_paragraph("")
    p = doc.add_paragraph()
    p.add_run("Interessenskonflikte").bold = True
    doc.add_paragraph("")
    doc.add_paragraph("")
    doc.add_page_break()

    img_n = tab_n = 0
    bangs = 1
    for b in blocks:
        if b.get("section"):
            bangs = 1  # neuer Hauptabschnitt: wieder ein Ausrufezeichen erlaubt
            sp = doc.add_paragraph()
            sp.add_run(f"Abschnitt: „{b['section'][0]}\t{b['section'][1]}“").bold = True
        if b.get("subsection"):
            up = doc.add_paragraph()
            up.add_run(f"Unterabschnitt: „{b['subsection'][0]}\t{b['subsection'][1]}“").bold = True
        if b.get("number"):
            mp = doc.add_paragraph()
            mp.add_run(f"Thema: „{b['number']}\t{b['topic']}“").bold = True
        elif not (b.get("section") or b.get("subsection")):
            mp = doc.add_paragraph()
            mp.add_run(f"Thema: „{b['topic']}“").bold = True

        sota = b.get("sota")
        if sota and sota.get("text"):
            doc.add_paragraph("")
            hp = doc.add_paragraph()
            hp.add_run("State of the Art").underline = True
            doc.add_paragraph(sota["text"].strip())
            if sota.get("literature"):
                doc.add_paragraph("")
                lp = doc.add_paragraph()
                lp.add_run(lit_label).bold = True
                for i, e in enumerate(sota["literature"], start=1):
                    doc.add_paragraph(f"{i}. {e}")
            doc.add_paragraph("")

        pubs = sorted(b["pubs"], key=lambda x: bool(x.get("kurz")))  # Kurznachrichten zuletzt
        kurz_started = False
        for pub in pubs:
            if pub.get("kurz") and not kurz_started:
                kurz_started = True
                if b.get("comment"):
                    _render_comment(doc, b["comment"], bangs, lang)
                    bangs = tame_exclamations(b["comment"], True, bangs)[1]
                doc.add_paragraph("")
                kp = doc.add_paragraph()
                kp.add_run(KURZ_HEADINGS["Englisch" if lang == "Englisch" else "Deutsch"]).underline = True
            title = (KURZ_MARK if pub.get("kurz") else "") + pub["title"]
            doc.add_heading(title, level=2).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if pub.get("recs"):  # gemeinsame Würdigung: alle Publikationen in der Zitierzeile
                doc.add_paragraph("; ".join(short_citation(r, [i]) for i, r in enumerate(pub["recs"], start=1)))
            else:
                doc.add_paragraph(short_citation(pub["rec"], [1]))
            if pub.get("abstract"):
                doc.add_paragraph(pub["abstract"])
            in_box = False
            in_comment = False
            for line in pub["lines"]:
                line = re.sub(r"^#+\s*", "", line.strip())
                if line and not re.fullmatch(r"[-*_—―]{3,}", line):
                    label = _label_of(line)
                    if label:
                        in_comment = label in _COMMENT_LABELS
                    line, bangs = tame_exclamations(line, in_comment, bangs)
                    in_box = render_line_with_bold_label(doc, line, in_box)
            if pub.get("literature"):
                doc.add_paragraph("")
                lp = doc.add_paragraph()
                lp.add_run(lit_label).bold = True
                for i, e in enumerate(pub["literature"], start=1):
                    doc.add_paragraph(f"{i}. {e}")
            fig = pub.get("figure")
            if fig and fig.get("bytes"):
                try:
                    ip = doc.add_paragraph()
                    add_floating_picture(ip, io.BytesIO(fig["bytes"]), Inches(5.5))
                    if fig.get("typ") == "Tabelle":
                        tab_n += 1
                        prefix = f"Tab. {tab_n}"
                    else:
                        img_n += 1
                        prefix = f"Abb. {img_n}"
                    authors = pub["rec"].get("authors") or ["?"]
                    ref = f"{authors[0]} et al., {pub['rec'].get('journal', '')}. {pub['rec'].get('year', '')}"
                    cp = doc.add_paragraph()
                    cp.add_run(f"{prefix}: {fig.get('caption', '').rstrip('.')}. {ref} (1)").italic = True
                except Exception:
                    pass
            doc.add_paragraph("")
        if b.get("comment") and not kurz_started:  # keine Kurznachrichten: Kommentar am Ende des Themas
            _render_comment(doc, b["comment"], bangs, lang)
            bangs = tame_exclamations(b["comment"], True, bangs)[1]

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_field(footer, "PAGE")
    footer.add_run("-")
    add_field(footer, "NUMPAGES")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _render_comment(doc: Document, text: str, bangs: int, lang: str) -> None:
    """Gemeinsamer Kommentar eines Themas als grauer Kasten mit fettem Label (wie der Kommentar einer
    einzelnen Würdigung). Ohne Literaturziffern, damit die Zusammenführung ihn nicht umnummeriert."""
    text = re.sub(r"\s*[(\[]\d+(?:\s*,\s*\d+)*[)\]]", "", text).strip()
    text, _ = tame_exclamations(text, True, bangs)
    label = "Comment: " if lang == "Englisch" else "Kommentar: "
    render_line_with_bold_label(doc, label + text, False)
    doc.add_paragraph("")


def build_recommendation_docx(chapter: str, topics: list[dict], records: dict, rankings: dict) -> bytes:
    """Liste der empfohlenen bzw. ausgewählten Publikationen (zum Beschaffen der Volltexte).
    topics: [{"title", "pmids"}] – jede Publikation steht nur einmal in der Liste."""
    doc = Document()
    style_document(doc)
    h = doc.add_paragraph()
    rr = h.add_run(f"Literaturempfehlung – {chapter}" if chapter else "Literaturempfehlung")
    rr.bold = True
    rr.font.size = Pt(14)
    doc.add_paragraph(date.today().strftime("%d.%m.%Y"))
    seen: set[str] = set()
    for t in topics:
        pmids = [p for p in dict.fromkeys(t.get("pmids", [])) if p not in seen]
        seen.update(pmids)
        if not pmids:
            continue
        doc.add_heading(t["title"], level=2)
        for n, pmid in enumerate(pmids, start=1):
            rec = records.get(pmid, {})
            it = rankings.get(pmid, {})
            authors = rec.get("authors") or []
            a = ", ".join(authors[:3]) + (" et al." if len(authors) > 3 else "")
            src = f"{rec.get('journal', '')}. {rec.get('year', '')}"
            if rec.get("volume"):
                src += f";{rec['volume']}" + (f"({rec['issue']})" if rec.get("issue") else "")
            if rec.get("pages"):
                src += f":{rec['pages']}"
            p = doc.add_paragraph()
            p.add_run(f"{n}. {a.rstrip('.')}. ")
            p.add_run(rec.get("title", "").rstrip(".") + ".").bold = True
            p.add_run(f" {src.strip()}." + (f" PMID: {pmid}" if pmid.isdigit() else ""))
            if rec.get("doi"):
                p.add_run(f" · doi:{rec['doi']}")
            if rec.get("is_oa"):
                p.add_run(" [Open Access]").italic = True
            if it:
                q = doc.add_paragraph()
                q.add_run(f"Priorität {it.get('priority', '?')} · Qualität {it.get('quality', '?')} · "
                          f"Bedeutung {it.get('question_importance', '?')} · "
                          f"IF {it.get('journal_impact_factor') or 'k. A.'} · "
                          f"{it.get('study_design', '')}. {it.get('rationale', '')}").italic = True
    footer = doc.sections[0].footer.paragraphs[0]  # Seitenzahlen
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_field(footer, "PAGE")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
