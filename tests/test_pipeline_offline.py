"""Offline-Test des Gesamtablaufs: KI, PubMed und Downloads werden simuliert.
Aufruf: python -m pytest tests/  (oder python tests/test_pipeline_offline.py)"""

import io
import json
import os
import re
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import fitz
from docx import Document

from medupdate import docx_tools, guidelines, manuscript, pipeline, pubmed, recommend, references

APPRAISAL = """1. Strengths of the Study
Randomized design with 1,200 patients (2).
2. Weaknesses and Limitations
Open-label design.
3. Examination of the Authors' Conclusions
Supported.
4. Context within the Current Evidence Base
Consistent with the guideline (3) and a meta-analysis (2).
5. Clinical Practice Implications
Relevant for adults.
Comment
A clear trial.
Conclusion for Practice
New, valid, practice changing for a subgroup.
References
1. Unused X. Not cited. J. 2020.
2. Meta A, Analysis B. A meta-analysis. Lancet. 2025;400:1-10. PMID: 99999999
3. German Society. S3 guideline. AWMF 2025."""

SHORT = """Study: A randomized trial in 1,200 patients.
Results: HR 0.80 (95% CI 0.70-0.90) (2).
Authors' Conclusion: Effective.
Conclusion for Practice: Practice changing for a subgroup (3)."""


class FakeLLM:
    def __init__(self):
        self.calls = []

    def text(self, prompt, **kw):
        p = prompt if isinstance(prompt, str) else json.dumps(prompt, default=str)
        self.calls.append(("text", p[:80]))
        if "Leitlinien zum folgenden Thema" in p:
            return "1. S3-Leitlinie Testthema, DGXX, 2025, https://register.awmf.org/x.pdf"
        if "ONE joint, comparative critical appraisal" in p:
            return ("Overview of the Publications\nPublication 1: RCT A.\nPublication 2: RCT B.\n"
                    "Assessment of the Linkage\nComparable populations, no pooled estimate.\n"
                    "Critical Appraisal\nBoth trials are consistent with the guideline (1).\n"
                    "Joint Conclusion\nBoth techniques are effective.\nReferences\n1. Joint Guideline X. Gut. 2025;1:1.")
        if "kürzt eine bereits erstellte gemeinsame" in p:
            return ("Publication 1: RCT A. Publication 2: RCT B. No pooled estimate. Consistent with the "
                    "guideline (1).\nJoint Conclusion: Both techniques are effective.")
        if "Du bist ein erfahrener klinisch-wissenschaftlicher Experte" in p:
            return APPRAISAL
        if "kürzt eine bereits erstellte kritische" in p:
            return SHORT
        if "gekürzte kritische Würdigung einer klinischen Studie vollständig ins" in p:
            return SHORT.replace("Study:", "Studie:").replace("Results:", "Ergebnisse:").replace(
                "Authors' Conclusion:", "Schlussfolgerung der Autoren:").replace(
                "Conclusion for Practice:", "Fazit für die Praxis:").replace("0.80", "0,80").replace(
                "0.70-0.90", "0,70-0,90")
        if "Einleitung der Publikation:" in p:
            assert "Knackpunkt" in p and "keine Stichpunkte" in p
            return "State of the Art\nThe burden is high (1). Standard care has limits (2).\nReferences\n1. Intro A. 2024.\n2. Intro B. Gut. 2023;1:1."
        if "State of the" in p and "Art" in p and "Baseline" in p:
            return "State of the Art\nThe field is established (1). Gaps remain (2).\nReferences\n1. Guideline A. 2025.\n2. Review B. Lancet. 2024;1:1."
        if "Übersetze den folgenden Absatz" in p:
            return "Das Fachgebiet ist etabliert (1). Lücken bleiben (2)."
        if "Schreibstil der Autorin" in p and "Analysiere" in p:
            return "- kurze Sätze"
        if "Passe den folgenden deutschen Text" in p:
            return p.split("Text:\n", 1)[1]
        return "OK"

    def text_ex(self, prompt, **kw):
        numbered = re.findall(r"^(\d+)\|\|\|(.*)$", prompt, re.MULTILINE)
        return "\n".join(f"{n}|||{t}" for n, t in numbered), False

    def json(self, prompt, schema, **kw):
        p = prompt if isinstance(prompt, str) else json.dumps(prompt, default=str)
        props = schema["properties"]
        if "groups" in props and "pub_ids" in json.dumps(schema):
            ids = re.findall(r"- id (\d+)", p)
            if "Thema Zwei" in p and len(ids) >= 2:
                return {"groups": [{"pub_ids": ids[:2], "question_de": "Zwei Techniken im Vergleich",
                                    "question_en": "Two techniques compared"}]}
            return {"groups": []}
        if "refs" in props:
            pid = re.search(r"- id (\S+)", p).group(1)
            return {"refs": [{"pub_id": pid, "manuscript": "alt.docx",
                              "sentence_de": "Im Jahr 2025 hatten wir dazu eine Vorstudie vorgestellt.",
                              "sentence_en": "In 2025 we presented a pilot study on this."}]}
        if "topics" in props:
            return {"chapter_number": "8", "chapter_title": "Gastroenterologie",
                    "topics": [{"number": "8.1.1", "title": "Thema Eins", "english_search_focus": "topic one",
                                "section_number": "8.1", "section_title": "Abschnitt"},
                               {"number": "8.1.2", "title": "Thema Zwei", "english_search_focus": "topic two",
                                "section_number": "8.1", "section_title": "Abschnitt"},
                               {"number": "8.2.1", "title": "Thema Drei", "english_search_focus": "topic three",
                                "section_number": "8.2", "section_title": "Leerer Abschnitt"},
                               {"number": "8.3.1", "title": "Thema Vier", "english_search_focus": "topic four",
                                "section_number": "8.3", "section_title": "Dritter Abschnitt"}]}
        if "citation_style" in props:
            return {k: (0 if v.get("type") == "integer" else ([] if v.get("type") == "array" else ""))
                    for k, v in props.items()} | {"max_pages": 20, "citation_style": "Vancouver", "recency_months": 12}
        if "queries" in props:
            return {"pico": "P", "queries": [{"name": "praezise", "query": "a"}, {"name": "breit", "query": "b"}],
                    "exclusion": ""}
        if "items" in props:
            ids = re.findall(r"PMID: (\d+)", p)
            # gerade n: Priorität A, ungerade n: B (geringe Praxisrelevanz); in "Thema Eins" fällt Fragestellung q1
            # wegen geringer Relevanz heraus, damit "Thema Zwei" mehr Publikationen erhält
            low = "Thema Eins" in p
            return {"items": [{"pmid": i, "study_design": "RCT", "relevance": 40 if (low and n % 3 == 1) else 90,
                               "evidence": 80 - n, "risk_of_bias": "Low", "novelty": 70,
                               "practice_impact": 80 if n % 2 == 0 else 50, "question_importance": 90 - 10 * n,
                               "journal_impact_factor": 50.0,
                               "question_key": f"q{n % 3}", "is_primary_research": True, "rationale": "gut"}
                              for n, i in enumerate(ids)]}
        if "guidelines" in props:
            return {"guidelines": [{"title": "S3-Leitlinie Test", "society": "DGXX", "year": "2025", "register_number": "",
                                    "page_url": "", "pdf_url": "https://x/y.pdf", "freely_available": True,
                                    "relevance": ""}]}
        if "entries" in props:
            ents = re.findall(r"^\d+\|\|\|(.*)$", p, re.MULTILINE)
            return {"entries": ents}
        if "index" in props:
            return {"index": 0, "typ": "Abbildung", "caption_en": "Kaplan-Meier curve", "caption_de": "Kaplan-Meier-Kurve"}
        if "assignments" in props:
            return {"assignments": []}
        if "befunde" in props:
            return {"befunde": []}
        if "kind" in props:
            return {"title": "Unknown", "authors": [], "journal": "", "year": "", "doi": "", "kind": "other"}
        raise AssertionError(f"unerwartetes Schema {list(props)}")


def _rec(pmid):
    return {"pmid": pmid, "title": f"Trial number {pmid} of something important", "abstract": "A" * 400 + " results.",
            "authors": ["Smith J", "Doe A", "Roe B", "Poe C"], "journal": "N Engl J Med", "journal_full": "",
            "year": "2026", "volume": "394", "issue": "2", "pages": "100-110", "doi": f"10.1056/x{pmid}",
            "pmc": "", "pub_types": ["Randomized Controlled Trial"], "language": "eng"}


def _pdf_with_image(text: str) -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), text)
    pix = fitz.Pixmap(fitz.csRGB, 400, 300, os.urandom(400 * 300 * 3), False)  # Rauschen: groß genug
    page.insert_image(fitz.Rect(50, 100, 450, 400), pixmap=pix)
    return doc.tobytes()


def _template() -> bytes:
    d = Document()
    for name in ("MainTitle1", "StandardTitle2", "Author", "Biblio", "BiblioTitle1", "ImageSubtitle", "Toc1", "Toc2"):
        d.styles.add_style(name, 1)
    d.add_paragraph("8\tGastroenterologie", style="MainTitle1")
    d.add_paragraph("Interessenkonflikte")
    d.add_paragraph("StandardTitle2 Beispiel", style="StandardTitle2")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def test_full_pipeline(monkeypatch=None):
    recommend.MAX_PER_QUESTION = 1  # feste Auswahl für die Zählungen unten; Gruppierung siehe Schritt 2
    counter = iter(range(10000000, 10001000))
    pubmed.search = lambda q, n, email="": ([str(next(counter)) for _ in range(6)], 6)
    pubmed.fetch_records = lambda ids, email="": {i: _rec(i) for i in ids}
    pubmed.find_pmid = lambda doi="", title="", email="": None
    pipeline.unpaywall = lambda doi, email: {"is_oa": False}
    pipeline.fetch_oa_pdf = lambda doi, email: None
    guidelines.download_pdf = lambda url: _pdf_with_image("Leitlinie Empfehlung " * 400)
    references.pubmed.find_pmid = lambda **kw: None

    llm = FakeLLM()
    proj = pipeline.new_project()
    proj["email"] = "test@example.org"
    log = lambda m: print("LOG", m)
    pipeline.run_intake(llm, proj, "8 Gastro\n8.1 Thema Eins\n8.2 Thema Zwei", "Hinweise", ("t.docx", _template()), log)
    pipeline.run_recommendations(llm, proj, log)
    pipeline.run_guideline_search(llm, proj, log)
    t1, t2 = proj["topics"]["T1"], proj["topics"]["T2"]
    assert len(t1["recommended"]) == 2 and len(t2["recommended"]) == 3, "nur Priorität A, eine je Fragestellung"
    assert all(t["ranking"][p]["priority"] == "A" for t in (t1, t2) for p in t["recommended"])
    assert len(set(proj["length_plan"]["tiers"].values())) > 1, "Länge richtet sich nach der Qualität"
    assert proj["topics"]["T1"]["guidelines"][0]["text"]
    pipeline.recommendation_docx(proj)

    proj["topics"]["T3"]["enabled"] = False  # ohne Publikationen: muss aus der Gliederung verschwinden
    first = proj["topics"]["T1"]["recommended"][0]
    rec = proj["records"][first]
    uploads = {"a.pdf": _pdf_with_image(f"{rec['title']} Smith J et al. " + "body text " * 500)}
    pipeline.match_uploads(llm, proj, uploads, log)
    second = proj["topics"]["T2"]["recommended"][0]
    pipeline.match_uploads(llm, proj, {"b.pdf": _pdf_with_image(
        f"{proj['records'][second]['title']} Smith J et al. " + "intro text " * 500)}, log)
    assert proj["uploads"]["a.pdf"]["pmid"] == first
    proj["style"]["alt.docx"] = "Früheres Manuskript. " * 300
    pipeline.run_manuscript(llm, proj, log, {"figures": True, "sota": True, "style": True, "trial": True,
                                             "trial_topic": "T2"})
    assert all(p in proj["topics"]["T2"]["recommended"] for p, d in proj["pubs"].items()
               if d.get("appraisal_en") and not d.get("group")), "Probelauf zum gewählten Thema"
    done = [p for p, d in proj["pubs"].items() if d.get("appraisal_en") and not d.get("group")]
    assert len(done) == 2, "Probelauf: höchstens 2 Publikationen"
    assert pipeline.partial_docx(proj)
    pipeline.match_uploads(llm, proj, {"c.pdf": _pdf_with_image("Knee arthroplasty outcomes in athletes. " * 300)}, log)
    assert "VARIA" in proj["topics"], "Upload ohne passendes Thema kommt unter Varia"
    pipeline.run_manuscript(llm, proj, log, {"figures": True, "sota": True, "style": True})
    assert not proj["errors"], proj["errors"]
    out = proj["outputs"]
    doc = Document(io.BytesIO(out["manuskript_DE"]))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Fazit für die Praxis" in text and "Studie:" in text
    assert text.count("Im Jahr 2025 hatten wir dazu eine Vorstudie vorgestellt.") == 1, "Rückbezug auf früheres Manuskript"
    assert "Gemeinsame Würdigung: Zwei Techniken im Vergleich" in text, "gemeinsame Würdigung"
    assert "Biblio" in {p.style.name for p in doc.paragraphs}
    biblio = [p.text for p in doc.paragraphs if p.style.name == "Biblio"]
    # 6 Publikationen + je Thema 2 SotA-Quellen + Meta-Analyse + Leitlinie
    print("\n".join(biblio))
    assert any("Smith J, Doe A, Roe B et al" in b for b in biblio)
    assert not any("Unused X" in b for b in biblio), "nicht zitierte Quelle darf nicht erscheinen"
    assert any("Joint Guideline X" in b for b in biblio), "externe Quelle der gemeinsamen Würdigung"
    gid = next(p for p, d in proj["pubs"].items() if d.get("group"))
    for m in proj["pubs"][gid]["group"]:
        assert any(proj["records"][m]["doi"] in b for b in biblio), "beide Publikationen im Literaturverzeichnis"
    raw = [p.text for p in Document(io.BytesIO(out["rohfassung_DE"])).paragraphs]
    assert proj["topics"]["T2"]["sota_en"]["source"] == "intro_v3", "State of the Art aus den Einleitungen"
    assert proj["topics"]["T4"]["sota_en"]["source"] == "web", "ohne Volltext: wie bisher per Websuche"
    assert raw.count("State of the Art") == 2, "genau ein State of the Art je Hauptabschnitt mit Publikationen"
    i_sota, i_t2 = raw.index("State of the Art"), raw.index("Thema: „8.1.2\tThema Zwei“")
    assert i_sota == i_t2 + 2, "State of the Art vor dem Thema mit den meisten Publikationen"
    paras = doc.paragraphs
    k = next(i for i, p in enumerate(paras) if p.text == "State of the Art")
    run = paras[k].runs[0]
    assert run.underline and not run.bold, "State of the Art unterstrichen, nicht fett"
    assert paras[k - 1].text == "" and paras[k + 1].text.strip(), "Leerzeile vor, nicht nach der Überschrift"
    toc = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name in ("Toc1", "Toc2")]
    print("GLIEDERUNG", toc)
    expected = [("Toc1", "8.1\tAbschnitt"), ("Toc2", "8.1.1\tThema Eins"), ("Toc2", "8.1.2\tThema Zwei"),
                ("Toc1", "8.2\tDritter Abschnitt"), ("Toc2", "8.2.1\tThema Vier"), ("Toc1", "8.3\tVaria")]
    assert toc[:6] == expected, "Inhaltsverzeichnis auf Seite 1, leerer Abschnitt entfernt und neu nummeriert"
    assert toc[6:] == expected, "gleiche Nummern als Überschriften im Text"
    assert not any("Leerer Abschnitt" in t or "Thema Drei" in t for _, t in toc)
    texts = [p.text for p in doc.paragraphs]
    assert texts.index("8.1\tAbschnitt") < texts.index("Interessenkonflikte"), "Inhaltsverzeichnis auf der ersten Seite"
    assert "Inhalt" not in texts and "Inhaltsverzeichnis" not in texts, "ohne Überschrift „Inhalt“"
    n_img = sum(1 for p in doc.paragraphs if p.text.startswith("Abb."))
    assert n_img == 0, "keine automatischen Abbildungen mehr"
    print("BEFUNDE")
    for b in out["befunde"]:
        print(b["status"], b["punkt"], "–", b["begruendung"][:100])
    old = next(p for p, d in proj["pubs"].items() if d.get("short_en"))
    proj["pubs"][old]["short_v"] = 1
    proj["pubs"][old]["short_en"] = "ALT"
    pipeline.run_manuscript(llm, proj, log, {"figures": True, "sota": True, "style": True, "reshorten": True})
    assert proj["pubs"][old]["short_en"] != "ALT" and proj["pubs"][old]["short_v"] == pipeline.SHORT_VERSION
    blob = pipeline.export_project(proj)
    again = pipeline.import_project(blob)
    assert again["outputs"]["manuskript_DE"] == out["manuskript_DE"]
    (Path(tempfile.gettempdir()) / "medupdate_test_DE.docx").write_bytes(out["manuskript_DE"])


def test_llm_retry_and_cost():
    import types

    import anthropic
    import httpx2 as httpx

    from medupdate import llm as L

    req = httpx.Request("POST", "https://api.anthropic.com/v1/messages")
    calls = {"n": 0}

    class Stream:
        def __enter__(self):
            calls["n"] += 1
            if calls["n"] == 1:
                raise anthropic.RateLimitError("rate", response=httpx.Response(429, request=req), body=None)
            return self

        def __exit__(self, *a):
            return False

        def get_final_message(self):
            usage = types.SimpleNamespace(input_tokens=1_000_000, output_tokens=100_000, cache_creation_input_tokens=0,
                                          cache_read_input_tokens=0,
                                          server_tool_use=types.SimpleNamespace(web_search_requests=3))
            return types.SimpleNamespace(content=[types.SimpleNamespace(type="text", text="ok")],
                                         stop_reason="end_turn", usage=usage, model="claude-opus-5-5")

    seen = {}

    def stream(**kw):
        seen.update(kw)
        return Stream()

    llm = L.LLM("x")
    llm.client = types.SimpleNamespace(beta=types.SimpleNamespace(messages=types.SimpleNamespace(stream=stream)))
    old_sleep, L.time.sleep = L.time.sleep, lambda s: None
    try:
        assert llm.text("hi", web_search=4) == "ok" and calls["n"] == 2
        assert seen["messages"][-1]["content"][0]["cache_control"] == {"type": "ephemeral"}
        llm.budget_usd = 5.0  # bereits ≈ 6 USD verbraucht → nächster Aufruf wird verweigert
        try:
            llm.text("hi")
            raise AssertionError("Kostenlimit muss greifen")
        except L.BudgetExceeded:
            pass
        llm.budget_usd = 0.0
    finally:
        L.time.sleep = old_sleep
    assert abs(llm.cost_usd - (4.0 + 2.0 + 0.03)) < 1e-6 and llm.web_searches == 3

    def billing(**kw):
        raise anthropic.BadRequestError("Your credit balance is too low", response=httpx.Response(400, request=req),
                                        body=None)
    llm.client = types.SimpleNamespace(beta=types.SimpleNamespace(messages=types.SimpleNamespace(stream=billing)))
    try:
        llm.text("hi")
        raise AssertionError("Kontofehler erwartet")
    except L.AccountError as ex:
        assert "Guthaben" in str(ex)


def test_results_saved_before_abort():
    from medupdate.llm import BudgetExceeded
    store = {}

    def fn(i):
        if i == 3:
            raise BudgetExceeded("Limit")
        return i * 10

    try:
        for k, v in pipeline._parallel([(i, i) for i in range(6)], fn, lambda m: None, "x", workers=1).items():
            store[k] = v
        raise AssertionError("Abbruch erwartet")
    except BudgetExceeded:
        pass
    assert store == {0: 0, 1: 10, 2: 20}, store


def test_search_broadening():
    from medupdate import recommend
    calls = []

    def fake_search(q, n, email=""):
        calls.append(q)
        return ([] if "narrow" in q else [str(20000000 + len(calls) * 10 + i) for i in range(5)]), 5

    class L:
        def json(self, prompt, schema, **kw):
            if "SENSITIVERE" in prompt:
                return {"pico": "", "queries": [{"name": "breit", "query": "wide"}], "exclusion": ""}
            return {"pico": "P", "queries": [{"name": "praezise", "query": "narrow"}], "exclusion": ""}

    old_search, old_fetch = pubmed.search, pubmed.fetch_records
    pubmed.search, pubmed.fetch_records = fake_search, (lambda ids, email="": {})
    try:
        res = recommend.recommend_for_topic(L(), {"title": "Fundusvarizen"}, "", {}, "", 6)
    finally:
        pubmed.search, pubmed.fetch_records = old_search, old_fetch
    assert len(res["attempts"]) == 3, "erweiterte Suche und längerer Zeitraum"
    assert res["window"][0] < recommend.time_window(12)[0].isoformat()

    pubmed.search = lambda q, n, email="": (_ for _ in ()).throw(RuntimeError("403"))
    try:
        recommend.collect_candidates({"queries": [{"name": "a", "query": "x"}]}, *recommend.time_window(12), "")
        raise AssertionError("Fehler muss sichtbar werden")
    except RuntimeError as ex:
        assert "PubMed" in str(ex)
    finally:
        pubmed.search = old_search


def test_length_budget():
    from medupdate.appraisal import plan_tiers
    pubs = [{"id": str(i), "length_score": s} for i, s in enumerate([90, 85, 70, 60, 50])]
    tiers, overflow = plan_tiers(pubs, {}, 1)
    assert tiers["0"] == "1,5 Seiten (Lang)" and tiers["2"] == "1 Seite" and tiers["4"] == "½ Seite (Kurz)"
    tiers, overflow = plan_tiers(pubs, {"max_pages": 4}, 1)
    assert overflow and overflow[0] == "4", "die schwächste Arbeit fällt zuerst heraus"
    assert all(t == "½ Seite (Kurz)" for t in tiers.values())
    tiers, overflow = plan_tiers(pubs, {"max_pages": 6}, 1)
    assert not overflow and tiers["0"] != "½ Seite (Kurz)" and tiers["4"] == "½ Seite (Kurz)"


def test_outline_levels():
    """Vierte Gliederungsebene: Hauptabschnitt > Unterabschnitt > Thema, neu nummeriert."""
    def t(title, sec, sub=""):
        return {"title": title, "section_number": sec, "section_title": f"Abschnitt {sec}",
                "subsection_number": sub, "subsection_title": f"Unter {sub}" if sub else ""}
    proj = {"toc": {"chapter_number": "8"}, "topics": {
        "T1": t("Direkt", "8.1"),
        "T2": t("Leer", "8.2", "8.2.1"),
        "T3": t("Varizen A", "8.2", "8.2.7"),
        "T4": t("Varizen B", "8.2", "8.2.7"),
        "T5": t("Einzeln", "8.2", "8.2.8"),
    }}
    g = pipeline.outline(proj, ["T1", "T3", "T4", "T5"])
    assert g["T1"] == {"section": ("8.1", "Abschnitt 8.1"), "subsection": None, "number": "8.1.1", "level": 2}
    assert g["T3"]["section"] == ("8.2", "Abschnitt 8.2") and g["T3"]["subsection"] == ("8.2.1", "Unter 8.2.7")
    assert g["T3"]["number"] == "8.2.1.1" and g["T4"]["number"] == "8.2.1.2" and g["T4"]["subsection"] is None
    assert g["T5"]["subsection"] == ("8.2.2", "Unter 8.2.8") and g["T5"]["number"] == "8.2.2.1"

    blocks = [{"topic": proj["topics"][tid]["title"], "sota": None, "pubs": [], **g[tid]}
              for tid in ("T1", "T3", "T4", "T5")]
    raw = Document(io.BytesIO(manuscript.build_docx("Titel", blocks, "Deutsch")))
    d = Document(io.BytesIO(_template()))
    d.styles.add_style("Toc3", 1)
    b = io.BytesIO()
    d.save(b)
    out = Document(io.BytesIO(docx_tools.copy_document_content_into_template(raw, b.getvalue())))
    toc = [(p.style.name, p.text) for p in out.paragraphs if p.style.name.startswith("Toc")]
    expected = [("Toc1", "8.1\tAbschnitt 8.1"), ("Toc2", "8.1.1\tDirekt"),
                ("Toc1", "8.2\tAbschnitt 8.2"), ("Toc2", "8.2.1\tUnter 8.2.7"),
                ("Toc3", "8.2.1.1\tVarizen A"), ("Toc3", "8.2.1.2\tVarizen B"),
                ("Toc2", "8.2.2\tUnter 8.2.8"), ("Toc3", "8.2.2.1\tEinzeln")]
    assert toc[:8] == expected, toc
    assert toc[8:] == expected, toc
    texts = [p.text for p in out.paragraphs]
    i = texts.index("8.1.1\tDirekt")
    assert texts[i + 1] == "" and texts[i + 2] == "8.2\tAbschnitt 8.2", "Leerzeile zwischen Hauptabschnitten"
    assert "Suchbegriffe" not in texts and not any(re.fullmatch(r"\d\d\.\d\d\.\d{4}", t) for t in texts)
    raw_texts = [p.text for p in raw.paragraphs]
    assert "Suchbegriffe" not in raw_texts and raw_texts[raw_texts.index("8.1.1\tDirekt") + 1] == ""


def test_sota_trailing_note():
    from medupdate.citations import split_sota_response
    raw = ("State of the Art\nA gilt (1). B gilt (2).\nLiteratur\n1. Pimentel-Nunes P. ESGE. Endoscopy. 2022.\n"
           "2. Ishihara R. Guidelines. Dig Endosc. 2020.\n\nNicht verifiziert:\n1. Dass A gilt.\n2. Dass B gilt. Bitte prüfen.")
    body, entries = split_sota_response(raw)
    assert entries == {1: "Pimentel-Nunes P. ESGE. Endoscopy. 2022.", 2: "Ishihara R. Guidelines. Dig Endosc. 2020."}, entries


def test_sota_choice():
    from medupdate.appraisal import sota_topics
    topics = {t: {"section_number": "8.1"} for t in ("A", "B", "C")} | {"D": {"section_number": "8.2"}}
    counts = {"A": 3, "B": 1, "C": 2, "D": 1}
    assert sota_topics(["A", "B", "C", "D"], topics, counts, {"A": 60, "B": 80, "C": 70, "D": 50}) == ["B", "D"]
    assert sota_topics(["A", "B", "C", "D"], topics, counts, {"A": 80, "B": 80, "C": 80}) == ["A", "D"]
    assert sota_topics(["A", "B", "C", "D"], topics, counts, {"A": 70, "B": 80, "C": 80}) == ["C", "D"]


def test_manual_window():
    from datetime import date
    assert recommend.manual_window(["2025-03-01", "2026-02-28"]) == (date(2025, 3, 1), date(2026, 2, 28))
    assert recommend.manual_window(None) is None
    seen = []
    orig = recommend.collect_candidates
    recommend.collect_candidates = lambda plan, von, bis, email: (seen.append((von, bis)) or ([], {}))
    try:
        res = recommend.recommend_for_topic(FakeLLM(), {"title": "Thema Eins"}, "Gastro", {}, "x@y.z", 6,
                                            window=["2025-03-01", "2026-02-28"])
    finally:
        recommend.collect_candidates = orig
    assert set(seen) == {(date(2025, 3, 1), date(2026, 2, 28))}, "manueller Zeitraum wird nicht verlängert"
    assert res["window"] == ["2025-03-01", "2026-02-28"]


def test_manual_limits():
    proj = {"rules": {"max_pages": 30, "max_references": 80}, "manual_limits": {"pages": 20, "references": 50}}
    pipeline.apply_manual_limits(proj)
    assert proj["rules"]["max_pages"] == 31 and proj["rules"]["max_references"] == 60, proj["rules"]
    proj["manual_limits"] = {"pages": 0, "references": 0}
    pipeline.apply_manual_limits(proj)
    assert proj["rules"]["max_pages"] == 30 and proj["rules"]["max_references"] == 80, "zurück zu den Hinweisen"
    proj["limit_tolerance"] = {"mode": "percent", "value": 10}
    proj["manual_limits"] = {"pages": 20, "references": 50}
    pipeline.apply_manual_limits(proj)
    assert proj["rules"]["max_pages"] == 23 and proj["rules"]["max_references"] == 55


def test_split_intro():
    from medupdate.appraisal import split_intro
    ft = ("Title\nAbstract\nShort.\nIntroduction\nVarices bleed often [1]. Therapy is standard [2].\n"
          "Methods\nWe did things.\n" + "x " * 500 + "\nReferences\n1. A B. Gut. 2020.\n2. C D. Lancet. 2021.")
    intro, refs = split_intro(ft)
    assert intro.startswith("Varices") and "Methods" not in intro and "We did" not in intro
    assert refs.startswith("1. A B")


def test_repeated_citations():
    from medupdate.citations import merge_repeated_citations
    d = Document()
    p = d.add_paragraph()
    p.add_run("Kommentar: ").bold = True
    p.add_run("Die Studie zeigt A (1). Im Jahr (2024) fand sich B (1,2). Der Effekt hielt an (1).")
    d.add_paragraph("Ergänzend gilt C (3). Die Daten stützen D (1).")
    d.add_paragraph("Studie: (1) Suche, (1) erneut")
    assert merge_repeated_citations(d) == 2
    assert d.paragraphs[0].text == "Kommentar: Die Studie zeigt A. Im Jahr (2024) fand sich B (2). Der Effekt hielt an (1).", \
        d.paragraphs[0].text
    assert d.paragraphs[1].text == "Ergänzend gilt C (3). Die Daten stützen D (1).", "neuer Absatz: wieder nennen"
    d2 = Document()
    d2.add_paragraph("A (1). B (1). C (4)")
    merge_repeated_citations(d2)
    assert d2.paragraphs[0].text == "A. B. C (1,4)", d2.paragraphs[0].text


def test_reassign_topics():
    class L:
        calls = 0

        def json(self, prompt, schema, **kw):
            L.calls += 1
            assert "Barrettkarzinom" in prompt and "[hochgeladen]" in prompt
            return {"assignments": [{"pub_id": "1", "topic_id": "T2", "reason_de": "Technischer Aspekt der Resektion."},
                                    {"pub_id": "2", "topic_id": "T1", "reason_de": ""},
                                    {"pub_id": "X", "topic_id": "T1", "reason_de": "Passt zu Barrett."}],
                    "same_question": [{"pub_ids": ["3", "2"], "question_de": "ESD versus EMR"}]}
    rk = {"question_key": "q", "length_score": 70}
    proj = {"toc": {"chapter_title": "GI"}, "records": {p: {"title": f"Studie {p}"} for p in ["1", "2", "3", "X"]},
            "topics": {"T1": {"id": "T1", "title": "Barrettkarzinom", "number": "1.1", "section_number": "1",
                              "section_title": "Ösophagus", "recommended": ["1", "2"], "ranking": {"1": rk, "2": rk}},
                       "T2": {"id": "T2", "title": "Endoskopische Resektion", "number": "1.2", "section_number": "1",
                              "section_title": "Ösophagus", "recommended": ["3"], "ranking": {"3": rk}}},
            "pubs": {"X": {"extra": True, "topic_id": "", "full_text": "Barrett text"}}}
    pipeline.reassign_topics(L(), proj, lambda m: None)
    cur = pipeline.current_topics(proj)
    # 1 technisch → Resektion; 2 und 3 gleiche Fragestellung → gemeinsames Thema (das von 3); X aus Varia → Barrett
    assert cur == {"1": "T2", "2": "T2", "3": "T2", "X": "T1"}, cur
    assert proj["topics"]["T2"]["ranking"]["1"] == rk
    assert proj["question_groups"] == [{"pmids": ["3", "2"], "question_de": "ESD versus EMR"}]
    ov = [r for r in pipeline.assignment_overview(proj) if r["kind"] == "topic"]
    t2 = {x["id"]: x for x in ov[1]["pubs"]}
    assert t2["1"]["moved_from"] == "Barrettkarzinom" and t2["2"]["joint"] == "ESD versus EMR"
    pipeline.reassign_topics(L(), proj, lambda m: None)
    assert L.calls == 1, "unveränderte Zuordnung wird nicht erneut geprüft"
    pipeline.set_manual_topic(proj, "1", "T1")
    assert pipeline.current_topics(proj)["1"] == "T1"
    pipeline.reassign_topics(L(), proj, lambda m: None)
    assert L.calls == 1 and pipeline.current_topics(proj)["1"] == "T1", "eigene Festlegung bleibt"
    assert pipeline.assignment_docx(proj)[:2] == b"PK"


def test_enforce_pub_limit():
    def it(q):
        return {"quality": q, "length_score": q, "evidence": q}
    proj = {"rules": {}, "manual_limits": {"publications": 3},
            "records": {p: {"title": f"Studie {p}"} for p in "abcde"},
            "topics": {"T1": {"id": "T1", "title": "Eins", "recommended": ["a", "b", "c"],
                              "ranking": {"a": it(90), "b": it(40), "c": it(70)}},
                       "T2": {"id": "T2", "title": "Zwei", "recommended": ["d"], "ranking": {"d": it(60), "e": it(10)}}},
            "pubs": {"e": {"extra": True, "topic_id": "T2"}}}
    dropped = pipeline.enforce_pub_limit(proj, lambda m: None)
    assert [e["id"] for e in dropped] == ["e", "b"], dropped
    assert set(pipeline.current_topics(proj)) == {"a", "c", "d"}
    assert "b" not in proj["topics"]["T1"].get("reserve", []) and proj["pubs"]["e"]["limit_dropped"]
    rev = proj["sel_rev"]
    proj["manual_limits"] = {"publications": 4}
    pipeline.enforce_pub_limit(proj, lambda m: None)
    assert set(pipeline.current_topics(proj)) == {"a", "b", "c", "d"} and proj["sel_rev"] > rev
    proj["manual_limits"] = {"publications": 0}
    assert pipeline.enforce_pub_limit(proj, lambda m: None) == [] and len(pipeline.current_topics(proj)) == 5
    # Seitenumfang: 4 Seiten inkl. Titelseite reichen nicht für alle fünf
    proj["manual_limits"] = {"pages": 3}
    proj["limit_tolerance"] = {"mode": "absolute", "value": 0}
    dropped = pipeline.enforce_pub_limit(proj, lambda m: None)
    assert dropped and dropped[0]["id"] == "e" and "Seitenumfang" in dropped[0]["reason"]


def test_direct_upload():
    import fitz
    def pdf(text):
        doc = fitz.open()
        doc.new_page().insert_text((50, 72), text, fontsize=9)
        return doc.tobytes()
    rec = {"title": "Endoscopic submucosal dissection versus EMR for Barrett neoplasia", "authors": ["Meier A"],
           "doi": "10.1000/xyz123"}
    proj = {"records": {"1": rec}, "pubs": {}, "uploads": {}}
    bad = pipeline.attach_direct_upload(proj, "1", "falsch.pdf", pdf("A completely different trial on liver cirrhosis outcomes"))
    assert not bad["ok"] and "liver" in bad["found"] and "1" not in proj["pubs"]
    good = pipeline.attach_direct_upload(proj, "1", "richtig.pdf", pdf("doi 10.1000/xyz123 Endoscopic submucosal dissection"))
    assert good["ok"] and proj["pubs"]["1"]["upload_name"] == "richtig.pdf" and proj["uploads"]["richtig.pdf"]["pmid"] == "1"
    forced = pipeline.attach_direct_upload(proj, "1", "falsch.pdf", pdf("Another unrelated paper about colon"), force=True)
    assert forced["ok"] and proj["pubs"]["1"]["upload_checked"] == "von Ihnen bestätigt"


def test_fazit_and_headings():
    class L:
        def json(self, prompt, schema, **kw):
            if "Überschriften" in prompt:
                ids = re.findall(r"^(\d+): (.+)$", prompt, re.M)
                return {"items": [{"id": i, "de": {"Esophagus": "Ösophagus", "Stomach": "Magen"}.get(x, x)}
                                  for i, x in ids]}
            ids = re.findall(r"^id (\d+):", prompt, re.M)
            return {"items": [{"id": i, "sentence": f"Variante {i} zeigt erstmals etwas."} for i in ids]}
    proj = {"toc": {"chapter_title": "Endoscopy"}, "pubs": {
        "a": {"short_de": "Studie: x\nFazit für die Praxis: Das Ergebnis ist neu. Es ist valide (1)."},
        "b": {"styled_de": "Kommentar: y\nFazit für die Praxis: Das Ergebnis ist neu. Praxis ändert sich nicht."},
        "c": {"short_de": "Fazit für die Praxis: Die Arbeit bestätigt frühere Daten."}},
        "topics": {"T": {"title": "Stomach", "section_title": "Esophagus"}}}
    n = pipeline.vary_fazit_openings(L(), proj, [("a", "T"), ("b", "T"), ("c", "T")], lambda m: None)
    assert n == 2
    assert proj["pubs"]["a"]["short_de"].endswith("Fazit für die Praxis: Variante 0 zeigt erstmals etwas. Es ist valide (1).")
    assert "Variante 1" in proj["pubs"]["b"]["styled_de"] and "bestätigt" in proj["pubs"]["c"]["short_de"]
    assert pipeline.vary_fazit_openings(L(), proj, [("a", "T"), ("b", "T")], lambda m: None) == 0
    de = pipeline.headings_de(L(), proj, ["T"])
    assert de["Esophagus"] == "Ösophagus" and de["Stomach"] == "Magen"


def test_comment_box():
    from docx import Document
    from medupdate.docx_tools import is_shaded, render_line_with_bold_label
    doc = Document()
    box = False
    for line in ["Studie: A.", "Kommentar: B.", "Weiter.", "Fazit für die Praxis: C."]:
        box = render_line_with_bold_label(doc, line, box)
    ps = doc.paragraphs
    assert not is_shaded(ps[0]) and ps[0].runs[0].underline, "Studie: unterstrichen, ohne Kasten"
    assert all(is_shaded(p) for p in ps[1:]), "Kommentar, Folgeabsatz und Fazit im grauen Kasten"
    assert ps[1].runs[0].bold and not ps[1].runs[0].underline, "Kommentar: fett, nicht unterstrichen"
    assert ps[3].runs[0].bold and not ps[3].runs[0].underline, "Fazit: fett, nicht unterstrichen"


def test_compact_note():
    from medupdate.appraisal import _compact_note
    assert _compact_note(None, "1 Seite") == ""
    n = _compact_note({"stats": True, "comment": True}, "½ Seite (Kurz)")
    assert "STATISTIK KOMPAKT" in n and "höchstens 1 kurze" in n and "KOMMENTAR KOMPAKT" in n
    assert "höchstens 3 kurze" in _compact_note({"comment": True}, "1,5 Seiten (Lang)")


if __name__ == "__main__":
    test_fazit_and_headings()
    test_direct_upload()
    test_enforce_pub_limit()
    test_reassign_topics()
    test_llm_retry_and_cost()
    test_results_saved_before_abort()
    test_search_broadening()
    test_length_budget()
    test_outline_levels()
    test_repeated_citations()
    test_split_intro()
    test_manual_limits()
    test_manual_window()
    test_sota_choice()
    test_sota_trailing_note()
    test_full_pipeline()
    test_comment_box()
    test_compact_note()
    print("OK")
